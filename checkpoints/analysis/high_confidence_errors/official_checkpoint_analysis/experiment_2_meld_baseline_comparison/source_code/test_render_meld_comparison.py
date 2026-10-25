"""Guard sparse-bin plots, CSV evidence, and whole-document image preservation."""
from pathlib import Path
import tempfile
import unittest

from docx import Document

from scripts.render_meld_comparison import (
    DOCX_NAME, FIGURE_STEM, make_figure, make_word, make_combined, append_document)
from scripts.render_public_m3net_comparison import verify_word
from scripts.render_official_reliability import _set_run, three_line_table
from tests.test_render_public_m3net_comparison import fixture_sources


def sources_fixture():
    sources = fixture_sources()
    for model_index, row in enumerate(sources['calibration_comparison.csv'][1]):
        row['model'] = ('Local README Baseline (test-selected)', 'Official CMERC')[model_index]
        row['valid_utterances'] = '2610'
    for row in sources['calibration_bins.csv'][1]:
        if row['model'] == 'Official M³Net':
            row['model'] = 'Local README Baseline (test-selected)'
        if row['count'] == '1623':
            row['count'] = '2610'
    sources['auxiliary_threshold_transitions.csv'] = sources['high_confidence_transitions.csv']
    return sources


class MeldReportTest(unittest.TestCase):
    def test_empty_bins_not_plotted_and_support_is_enforced(self):
        sources = sources_fixture()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            (path / 'figures').mkdir()
            result = make_figure(path, sources['calibration_bins.csv'][1])
            self.assertTrue(result['empty_bins_omitted'])
            self.assertEqual(result['model_series'][0]['nonempty_bin_indices'], [9])
            sources['calibration_bins.csv'][1][9]['count'] = '2609'
            with self.assertRaisesRegex(ValueError, '2610'):
                make_figure(path, sources['calibration_bins.csv'][1])

    def test_word_values_match_csv_and_numeric_tamper_is_rejected(self):
        sources = sources_fixture()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            (path / 'figures').mkdir()
            make_figure(path, sources['calibration_bins.csv'][1])
            records, headings = make_word(path, sources, 'Historical test selection is disclosed.')
            result = verify_word(path / DOCX_NAME, sources, records, headings, path / f'{FIGURE_STEM}.png')
            self.assertTrue(result['passed'])
            document = Document(path / DOCX_NAME)
            document.tables[0].rows[1].cells[1].text = '999'
            document.save(path / DOCX_NAME)
            with self.assertRaisesRegex(AssertionError, 'differs from source CSV'):
                verify_word(path / DOCX_NAME, sources, records, headings, path / f'{FIGURE_STEM}.png')

    def test_merge_retains_original_body_tables_and_multiple_images(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            (path / 'figures').mkdir()
            make_figure(path, sources_fixture()['calibration_bins.csv'][1])
            image = path / f'{FIGURE_STEM}.png'
            paths = []
            for index in range(3):
                document = Document()
                _set_run(document.add_paragraph().add_run(f'报告 {index} 原内容'), bold=True)
                three_line_table(document, ['计数', '结论'], [[str(index), '保留']])
                document.add_picture(str(image))
                for paragraph in document.paragraphs:
                    for run in paragraph.runs:
                        _set_run(run)
                source = path / f'source{index}.docx'
                document.save(source)
                paths.append(source)
            result = make_combined(*paths, path / 'combined.docx')
            self.assertTrue(result['original_body_prefix_identical'])
            self.assertEqual(result['original_tables_preserved'], 1)
            self.assertEqual(len(result['all_embedded_image_sha256']), 3)
            self.assertEqual(len(Document(path / 'combined.docx').inline_shapes), 3)


if __name__ == '__main__':
    unittest.main()
