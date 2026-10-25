"""Render the fixed public-weight comparison from checked CSVs, without inference.

The table provenance and saved OOXML are checked cell by cell. Existing outputs
are never overwritten. Figure points are the observed nonempty calibration bins.
"""
import argparse
import json
import math
import os
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET
import zipfile

import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

try:
    from .render_official_reliability import (
        _set_run, three_line_table, read_csv, digest, FONT, NS)
except ImportError:
    from render_official_reliability import (
        _set_run, three_line_table, read_csv, digest, FONT, NS)

DOCX_NAME = 'CMERC_官方权重_实验二分类纠错与校准对比.docx'
FIGURE_STEM = 'figures/reliability_IEMOCAP_public_comparison'
SOURCES = ('paired_predictions.csv', 'classification_transitions.csv',
    'high_confidence_transitions.csv', 'classwise_transitions.csv',
    'calibration_comparison.csv', 'calibration_bins.csv', 'typical_cases.csv')
EMOTIONS = ('happy', 'sad', 'neutral', 'angry', 'excited', 'frustrated')


def display(value, *, percent=False, integer=False, signed=False):
    if value in ('', 'N/A', None, 'None', 'null'):
        return 'N/A'
    number = float(value)
    if not math.isfinite(number):
        raise ValueError('Nonfinite report statistic')
    if integer:
        if not number.is_integer():
            raise ValueError('Expected an integer count')
        return str(int(number))
    return format(number * (100 if percent else 1),
                  '+.4f' if signed else '.4f' if percent else '.6f')


def figure_contract():
    return '''# Public M³Net–CMERC reliability comparison figure contract

Backend: the existing Python / matplotlib workflow exclusively.
Archetype: quantitative grid; fixed public-checkpoint descriptive comparison.
Claim: show each model's confidence–accuracy deviation and bin support under
one identical IEMOCAP test protocol, without attributing it to training modules.

Upper panel: empirical accuracy versus mean maximum probability for each
nonempty bin; identical [0,1] axes and ideal diagonal. Connected points only
guide the eye and are not fitted calibration curves. Empty bins have no point.
Lower panel: all ten equal-width-bin supports, grouped by fixed model.

Source: calibration_bins.csv; n=1623 valid utterances per model. No new model
inference, smoothing, relabeling, probability clipping or recalibration.
Exports: PNG at 300 dpi for Word; editable-text SVG; TrueType PDF;
6.2 × 6.8 inches, white background, two distinguishable colors and markers.
Risks: sparse-bin uncertainty is not estimated; no confidence intervals,
seed-level inference, significance or strict training-pair / causal claims.
'''


def make_figure(directory, rows):
    models = list(dict.fromkeys(row['model'] for row in rows))
    if len(models) != 2:
        raise ValueError('Expected exactly two fixed models')
    records = []
    with plt.rc_context({'font.family': 'DejaVu Sans', 'font.size': 9,
                         'svg.fonttype': 'none', 'pdf.fonttype': 42,
                         'axes.spines.top': False, 'axes.spines.right': False}):
        fig, (axis, support) = plt.subplots(2, 1, figsize=(6.2, 6.8),
            gridspec_kw={'height_ratios': [3.4, 1]})
        fig.subplots_adjust(left=.15, right=.97, top=.94, bottom=.09, hspace=.40)
        axis.plot([0, 1], [0, 1], '--', color='#777777', linewidth=1,
                  label='Ideal calibration')
        for index, (model, color, marker) in enumerate(zip(
                models, ('#326A98', '#AF6B38'), ('o', 's'))):
            bins = sorted([row for row in rows if row['model'] == model],
                          key=lambda row: int(row['bin_index']))
            if [int(r['bin_index']) for r in bins] != list(range(10)):
                raise ValueError('Missing/duplicate equal-width calibration bins')
            if sum(int(r['count']) for r in bins) != 1623:
                raise ValueError('Bin support must total 1623 per fixed model')
            active = [row for row in bins if int(row['count']) > 0]
            points = [[float(r['mean_confidence']), float(r['accuracy'])] for r in active]
            if any(not all(math.isfinite(v) and 0 <= v <= 1 for v in p) for p in points):
                raise ValueError('Invalid reliability point')
            axis.plot([p[0] for p in points], [p[1] for p in points],
                      color=color, marker=marker, markersize=4, linewidth=1.2, label=model)
            support.bar([(float(r['lower_bound']) + float(r['upper_bound'])) / 2
                         + (-.021 if index == 0 else .021) for r in bins],
                        [int(r['count']) for r in bins], width=.038,
                        color=color, label=model)
            records.append(dict(model=model, counts=[int(r['count']) for r in bins],
                nonempty_bin_indices=[int(r['bin_index']) for r in active], points=points))
        axis.set(xlim=(0, 1), ylim=(0, 1), xlabel='Mean confidence',
                 ylabel='Empirical accuracy', title='IEMOCAP: fixed public checkpoints')
        axis.set_aspect('equal', adjustable='box')
        axis.set_xticks([0, .2, .4, .6, .8, 1])
        axis.set_yticks([0, .2, .4, .6, .8, 1])
        axis.legend(loc='upper left', frameon=False, fontsize=8)
        support.set(xlim=(0, 1), xlabel='Confidence bin', ylabel='Utterances')
        support.set_xticks([0, .2, .4, .6, .8, 1])
        for extension in ('png', 'svg', 'pdf'):
            fig.savefig(directory / f'{FIGURE_STEM}.{extension}', dpi=300,
                        facecolor='white')
        plt.close(fig)
    return dict(source_csv='calibration_bins.csv', ideal_diagonal=True,
                empty_bins_omitted=True, model_series=records)


def make_word(directory, sources):
    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Mm(297), Mm(210)
        section.top_margin = section.bottom_margin = Mm(18)
        section.left_margin = section.right_margin = Mm(18)
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Caption'):
        style = document.styles[name]
        style.font.name, style.font.size = FONT, Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn(f'w:{slot}'), FONT)
    tables, heading_texts = [], []

    def para(text, heading=False):
        paragraph = document.add_paragraph()
        _set_run(paragraph.add_run(text), bold=heading)
        if heading:
            heading_texts.append(text)
        return paragraph

    def table(title, source, columns, headers, *, percents=(), integers=(), signed=(), indices=None):
        fields, rows = sources[source]
        if set(columns) - set(fields):
            raise ValueError(f'Missing columns in {source}: {set(columns) - set(fields)}')
        chosen = list(range(len(rows))) if indices is None else indices
        rendered = []
        for i in chosen:
            rendered.append([display(rows[i][field], percent=field in percents,
                integer=field in integers, signed=field in signed)
                if field in set(percents) | set(integers) | set(signed) or
                field in {'brier_score', 'nll'} or field.startswith(('baseline_probability_', 'cmerc_probability_'))
                else rows[i][field] for field in columns])
        para(title, True)
        record = dict(table_index=len(document.tables), source_csv=source,
            columns=columns, source_row_indices=chosen, headers=headers,
            percentage_columns=list(percents), integer_columns=list(integers),
            signed_columns=list(signed), rendered_rows=rendered)
        three_line_table(document, headers, rendered)
        tables.append(record)

    para('CMERC 官方权重：实验二分类纠错与概率校准对比', True)
    para('本报告对照 M³Net 与 CMERC 作者分别公开、预先固定的 IEMOCAP 权重。'
         '全部统计复用已保存的 1623 条完整文本、音频、视觉输入预测；'
         '不训练、不重新推理、不更换 checkpoint、不根据测试结果选模。')
    para('一、来源、评估协议与解释边界', True)
    para('M³Net 固定权重为 Roberta_IEMOCAP3.pth.tar；CMERC 固定权重为 IEMOCAP_checkpoint.pkl。'
         '两者模型结构、特征文件、六分类映射、测试话语、batch size=16、Padding/Mask、'
         '默认 cuDNN RNN 后端与已记录 PyG 兼容补丁已核验。'
         '本轮按 (dataset, dialogue_id, utterance_index) 重新核验 source utterance ID、原文和真实标签。')
    para('这是一组公开固定权重的描述性行为对照；训练 seed、Epoch、checkpoint 选择方式与'
         '两次训练的配对关系尚未确认。因此不构成严格模块消融，不能将差异归因于 CL、HSCL、CC '
         '或声称这些模块产生了因果收益。输入权重和预测文件 SHA-256 见 run_manifest.json。')
    para('二、逐样本分类变化', True)
    table('表 1　全测试集分类转移（比例分母为 1623）', 'classification_transitions.csv',
          ['transition', 'count', 'ratio'], ['分类转移', '话语数', '占全部话语 (%)'],
          percents=('ratio',), integers=('count',))
    comparison = sources['calibration_comparison.csv'][1]
    base, cm = comparison
    para(f'固定 M³Net 的错误数为 {base["total_errors"]}，固定 CMERC 的错误数为 {cm["total_errors"]}。'
         f'W-F1 分别为 {display(base["weighted_f1"], percent=True)}% 与 '
         f'{display(cm["weighted_f1"], percent=True)}%，CMERC−M³Net 为 '
         f'{display(float(cm["weighted_f1"])-float(base["weighted_f1"]), percent=True, signed=True)} 个百分点。'
         '分类纠正和新增错误必须同时报告；置信度降低本身不表示类别已纠正。')
    para('三、高置信度错误与逐阈值转移', True)
    table('表 2　高置信度预测的支持数与错误率', 'high_confidence_transitions.csv',
          ['threshold', 'baseline_high_confidence_predictions', 'baseline_high_confidence_errors',
           'baseline_high_confidence_error_rate', 'cmerc_high_confidence_predictions',
           'cmerc_high_confidence_errors', 'cmerc_high_confidence_error_rate'],
          ['阈值', 'M³Net 高置信预测', 'M³Net 高置信错误', 'M³Net 错误率 (%)',
           'CMERC 高置信预测', 'CMERC 高置信错误', 'CMERC 错误率 (%)'],
          percents=('baseline_high_confidence_error_rate', 'cmerc_high_confidence_error_rate'),
          integers=('baseline_high_confidence_predictions', 'baseline_high_confidence_errors',
                    'cmerc_high_confidence_predictions', 'cmerc_high_confidence_errors'))
    table('表 3　原 M³Net 高置信度错误的去向与 CMERC 新增错误',
          'high_confidence_transitions.csv',
          ['threshold', 'baseline_high_confidence_errors', 'corrected', 'still_wrong_confidence_lower',
           'still_wrong_confidence_not_lower', 'still_wrong_above_threshold', 'new_cmerc_hc_errors'],
          ['阈值', '原高置信错误', '分类纠正', '仍错且置信度下降', '仍错且置信度未降',
           '仍错且达到原阈值', 'CMERC 新增高置信错误'],
          integers=('baseline_high_confidence_errors', 'corrected', 'still_wrong_confidence_lower',
                    'still_wrong_confidence_not_lower', 'still_wrong_above_threshold', 'new_cmerc_hc_errors'))
    para('表注：阈值均为 ≥0.80、≥0.90、≥0.95；高置信度预测错误率的分母为相应模型所有'
         '达到阈值的预测数，分母为零时为 N/A。表 3 前三种去向互斥且合计等于原高置信错误数；'
         '“仍错且达到原阈值”是另外的重叠子集，不应与前三项相加。'
         '新增高置信错误指 CMERC 为高置信错误而 M³Net 原本并非相同阈值高置信错误，'
         '可包括 M³Net 原本正确或原本低置信错误的样本。')
    table('表 3a　原高置信度错误去向的比例', 'high_confidence_transitions.csv',
          ['threshold', 'corrected_ratio', 'still_wrong_confidence_lower_ratio',
           'still_wrong_confidence_not_lower_ratio', 'still_wrong_above_threshold_ratio'],
          ['阈值', '分类纠正 (%)', '仍错且置信度下降 (%)', '仍错且置信度未降 (%)',
           '仍错且达到原阈值 (%)'],
          percents=('corrected_ratio', 'still_wrong_confidence_lower_ratio',
                    'still_wrong_confidence_not_lower_ratio', 'still_wrong_above_threshold_ratio'))
    table('表 3b　CMERC 新增高置信错误的分母与来源', 'high_confidence_transitions.csv',
          ['threshold', 'eligible_non_baseline_hc_errors', 'new_cmerc_hc_errors',
           'new_cmerc_hc_error_ratio_of_eligible', 'new_cmerc_hc_errors_from_baseline_correct',
           'new_cmerc_hc_errors_from_baseline_wrong_below_threshold'],
          ['阈值', '原非高置信错误数', '新增高置信错误', '新增占可进入样本 (%)',
           '来自原正确', '来自原低置信错误'],
          percents=('new_cmerc_hc_error_ratio_of_eligible',),
          integers=('eligible_non_baseline_hc_errors', 'new_cmerc_hc_errors',
                    'new_cmerc_hc_errors_from_baseline_correct',
                    'new_cmerc_hc_errors_from_baseline_wrong_below_threshold'))
    para('表注：表 3a 的分母均为同阈值下 M³Net 的高置信度错误数；表 3b 的新增比例分母为'
         '1623−M³Net 高置信度错误数。新增来源两项互斥，合计为全部新增高置信错误数。')
    para('最大预测置信度变化与真实类别概率变化分别保存在 paired_predictions.csv。'
         '两个模型的错误预测类别可能不同，因此最大概率下降并不意味着同一错误类别概率下降；'
         '更不能单凭该变化认定存在模态冲突。')
    para('四、类别级纠错与退化', True)
    table('表 4　按真实情绪统计分类转移', 'classwise_transitions.csv',
          ['true_emotion', 'support', 'baseline_errors', 'cmerc_errors',
           'corrected', 'regressed', 'both_wrong', 'net_error_change'],
          ['真实情绪', '支持数', 'M³Net 错误', 'CMERC 错误', '错→对', '对→错', '两者都错', '净错误变化'],
          integers=('support', 'baseline_errors', 'cmerc_errors', 'corrected',
                    'regressed', 'both_wrong', 'net_error_change'))
    para('表注：净错误变化=CMERC 错误数−M³Net 错误数，正值表示错误增加。'
         '类别表使用真实情绪分组，支持数合计 1623；不按预测情绪分组替代。')
    para('五、全测试集概率校准对比', True)
    table('表 5　分类性能与错误置信度', 'calibration_comparison.csv',
          ['model', 'valid_utterances', 'total_errors', 'accuracy', 'weighted_f1',
           'mean_confidence', 'mean_error_confidence'],
          ['模型', '话语数', '错误数', 'Accuracy (%)', 'W-F1 (%)', '平均置信度 (%)', '错误平均置信度 (%)'],
          percents=('accuracy', 'weighted_f1', 'mean_confidence', 'mean_error_confidence'),
          integers=('valid_utterances', 'total_errors'))
    table('表 6　校准与概率评分', 'calibration_comparison.csv',
          ['model', 'ece', 'mce', 'brier_score', 'nll'],
          ['模型', 'ECE (%)', 'MCE (%)', 'Brier Score', 'NLL'], percents=('ece', 'mce'))
    para('表注：所有指标使用全部 1623 条有效话语。ECE/MCE 采用十个等宽箱 [0,0.1)、…、'
         '[0.9,1]，空箱不计入 MCE 最大值；ECE=Σ(n箱/N)|平均置信度−箱内准确率|。'
         'Brier 为所有类别平方误差之和的样本均值，不除以类别数或 2；'
         'NLL 为真实类别概率的负自然对数均值，不裁剪概率。分箱差值为平均置信度−准确率。')
    for model in dict.fromkeys(row['model'] for row in sources['calibration_bins.csv'][1]):
        indices = [i for i, row in enumerate(sources['calibration_bins.csv'][1]) if row['model'] == model]
        table(f'表　{model} 十等宽校准分箱', 'calibration_bins.csv',
              ['lower_bound', 'upper_bound', 'count', 'mean_confidence', 'accuracy', 'signed_gap'],
              ['下界', '上界', '话语数', '平均置信度 (%)', '实际准确率 (%)', '有符号差值 (pp)'],
              percents=('mean_confidence', 'accuracy', 'signed_gap'), integers=('count',),
              signed=('signed_gap',), indices=indices)
    document.add_picture(str(directory / f'{FIGURE_STEM}.png'), width=Inches(5.6))
    para('图 1　IEMOCAP 两份固定公开权重的可靠性对照。虚线为理想校准；各模型只绘制非空箱'
         '的真实均值点，连线仅引导阅读，不表示拟合；下方为全部十箱的样本支持数。'
         '空箱没有虚假观测点，稀疏箱不代表稳定的跨运行规律。')
    para('六、典型案例：纠正、降置信度与新增错误均展示', True)
    case_rows = sources['typical_cases.csv'][1]
    case_types = list(dict.fromkeys(row['case_type'] for row in case_rows))
    para('案例选择规则：typical_cases.csv 完整保留四种案例类型在三个阈值下的所有匹配样本；'
         '同一句可因类型或阈值重叠出现。正文统一展示阈值 0.90 下每类按预先定义的 selection_score '
         '降序、对话 ID 与话语索引稳定排序的前两条。概率表同时列两模型全部六类概率，'
         '不只选择支持改善的样本。完整逐话语证据见 paired_predictions.csv。')
    para('selection_score 的固定定义：纠正与仍错降置信度案例使用 M³Net 最大置信度；'
         '原正确而 CMERC 新增高置信错误使用 CMERC 最大置信度；两者均高置信错误使用两模型'
         '最大置信度的较小者。该排序只决定报告摘录，不用于选择 checkpoint 或修改统计。')
    for case_type in case_types:
        selected = [i for i, row in enumerate(case_rows)
                    if row['case_type'] == case_type and float(row['threshold']) == .9]
        selected = sorted(selected, key=lambda i: (-float(case_rows[i]['selection_score']),
                          case_rows[i]['dialogue_id'], int(case_rows[i]['utterance_index'])))[:2]
        para(f'案例类型：{case_type}', True)
        if not selected:
            para('该类型在 0.90 阈值下无匹配样本。')
        for i in selected:
            row = case_rows[i]
            para(f'对话 {row["dialogue_id"]}，话语索引 {row["utterance_index"]}，'
                 f'源 ID {row["source_utterance_id"]}，真实情绪 {row["true_emotion"]}。原文：{row["utterance_text"]}')
            table('案例预测、最大置信度与真实类别概率', 'typical_cases.csv',
                  ['baseline_predicted_emotion', 'baseline_predicted_confidence', 'baseline_true_label_probability',
                   'cmerc_predicted_emotion', 'cmerc_predicted_confidence', 'cmerc_true_label_probability'],
                  ['M³Net 预测', 'M³Net 最大概率 (%)', 'M³Net 真实类概率 (%)',
                   'CMERC 预测', 'CMERC 最大概率 (%)', 'CMERC 真实类概率 (%)'],
                  percents=('baseline_predicted_confidence', 'baseline_true_label_probability',
                            'cmerc_predicted_confidence', 'cmerc_true_label_probability'), indices=[i])
            for prefix, label in (('baseline', 'M³Net'), ('cmerc', 'CMERC')):
                table(f'{label} 六类概率（0–1）', 'typical_cases.csv',
                      [f'{prefix}_probability_{emotion}' for emotion in EMOTIONS],
                      list(EMOTIONS), indices=[i])
    para('七、结论与局限', True)
    para('这两份固定权重在分类与概率评分上表现不同：CMERC 的分类错误更多、W-F1 较低；'
         '其 ECE、MCE、Brier 与 NLL 的具体差值见表 6。较低的校准或概率评分不能替代分类纠错证据，'
         '也不意味着所有样本的置信度均得到改善。')
    para('CMERC−M³Net 的指标差值：'
         f'ECE {display(float(cm["ece"])-float(base["ece"]), percent=True, signed=True)} 个百分点；'
         f'MCE {display(float(cm["mce"])-float(base["mce"]), percent=True, signed=True)} 个百分点；'
         f'Brier {display(float(cm["brier_score"])-float(base["brier_score"]), signed=True)}；'
         f'NLL {display(float(cm["nll"])-float(base["nll"]), signed=True)}。'
         '这些数值的小幅改善与较低的分类 F1 同时存在，不能据此声称整体纠错能力提升。')
    para('全部结果是单个固定测试集上的描述性统计，未估计训练随机性、显著性或因果效果；'
         '训练历史和严格配对关系仍未知。后续若研究模块净收益，应使用预先固定的共同训练与'
         '验证选模协议进行配对实验；本任务不执行此类训练。')
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            _set_run(run)
    document.save(directory / DOCX_NAME)
    return tables, heading_texts


def verify_word(path, sources, records, headings, figure_path):
    document = Document(path)
    if len(document.tables) != len(records):
        raise AssertionError('Unexpected report table count')
    checked = 0
    for record in records:
        # Re-read the original CSV rather than trusting the renderer's stored rows.
        _, original_rows = sources[record['source_csv']]
        expected = []
        for i in record['source_row_indices']:
            row = []
            for field in record['columns']:
                value = original_rows[i][field]
                if field in set(record['percentage_columns']) | set(record['integer_columns']) | set(record['signed_columns']) or field in {'brier_score', 'nll'} or field.startswith(('baseline_probability_', 'cmerc_probability_')):
                    value = display(value, percent=field in record['percentage_columns'],
                        integer=field in record['integer_columns'], signed=field in record['signed_columns'])
                row.append(value)
            expected.append(row)
        actual = [[cell.text for cell in row.cells]
                  for row in document.tables[record['table_index']].rows]
        if actual != [record['headers']] + expected:
            raise AssertionError('Word table value differs from source CSV')
        checked += sum(map(len, actual))
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
        if any(name.startswith(('word/header', 'word/footer')) for name in archive.namelist()):
            raise AssertionError('Unexpected header/footer files')
        if root.findall('.//w:headerReference', NS) or root.findall('.//w:footerReference', NS):
            raise AssertionError('Unexpected header/footer reference')
        for run in root.findall('.//w:r', NS):
            fonts, color, size = (run.find('w:rPr/' + name, NS) for name in ('w:rFonts', 'w:color', 'w:sz'))
            if (fonts is None or any(fonts.get(qn(f'w:{slot}')) != FONT for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'))
                    or color is None or color.get(qn('w:val')) != '000000'
                    or size is None or size.get(qn('w:val')) != '24'):
                raise AssertionError('Incorrect explicit Song font, 12pt size or black color')
        for paragraph in root.findall('w:body/w:p', NS):
            text = ''.join(paragraph.itertext())
            if text in headings and any(run.find('w:rPr/w:b', NS) is None for run in paragraph.findall('w:r', NS)):
                raise AssertionError('Heading is not bold')
        for table in root.findall('.//w:tbl', NS):
            borders = table.find('w:tblPr/w:tblBorders', NS)
            if borders is None:
                raise AssertionError('Missing three-line table borders')
            for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
                edge = borders.find(f'w:{side}', NS)
                if edge is None or edge.get(qn('w:val')) != ('single' if side in ('top', 'bottom') else 'nil'):
                    raise AssertionError('Invalid table border')
            for cell in table.findall('w:tr', NS)[0].findall('w:tc', NS):
                edge = cell.find('w:tcPr/w:tcBorders/w:bottom', NS)
                if edge is None or edge.get(qn('w:val')) != 'single':
                    raise AssertionError('Missing rule below table header')
                if any(run.find('w:rPr/w:b', NS) is None for run in cell.findall('.//w:r', NS)):
                    raise AssertionError('Table header is not bold')
        images = [name for name in archive.namelist() if name.startswith('word/media/')]
        if len(images) != 1 or archive.read(images[0]) != Path(figure_path).read_bytes():
            raise AssertionError('Embedded reliability image differs from checked figure')
        styles = ET.fromstring(archive.read('word/styles.xml'))
        for name in ('Normal', 'Heading1', 'Heading2', 'Caption'):
            style = next(s for s in styles.findall('w:style', NS) if s.get(qn('w:styleId')) == name)
            fonts = style.find('w:rPr/w:rFonts', NS)
            if fonts is None or any(fonts.get(qn(f'w:{slot}')) != FONT for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs')):
                raise AssertionError('Report style font mismatch')
    return dict(passed=True, tables_checked=len(records), table_cells_checked=checked,
        all_table_cells_match_csv=True, heading_and_header_bold=True, font=FONT,
        font_size_points=12, color='000000', explicit_three_line_tables=True,
        headers_footers=False, embedded_figure_sha256=digest(figure_path),
        office_layout_rendered=False,
        layout_note='OOXML checked; Office/LibreOffice pagination was not rendered.',
        table_provenance=records)


def render(output):
    output = Path(output).resolve()
    sources = {name: read_csv(output / name) for name in SOURCES}
    hashes = {name: digest(output / name) for name in SOURCES}
    paths = [DOCX_NAME, 'report_verification.json', 'figure_contract.md'] + [f'{FIGURE_STEM}.{ext}' for ext in ('png', 'svg', 'pdf')]
    for name in paths:
        if (output / name).exists():
            raise FileExistsError(f'Refusing to overwrite report artifact: {output / name}')
    with tempfile.TemporaryDirectory(prefix='.public_render_', dir=output) as folder:
        temporary = Path(folder)
        (temporary / 'figures').mkdir()
        (temporary / 'figure_contract.md').write_text(figure_contract(), encoding='utf-8')
        figure = make_figure(temporary, sources['calibration_bins.csv'][1])
        records, headings = make_word(temporary, sources)
        result = dict(status='passed', training=False, gpu_inference=False,
            source_csv_sha256=hashes, figure=figure,
            word=verify_word(temporary / DOCX_NAME, sources, records, headings,
                             temporary / f'{FIGURE_STEM}.png'))
        if hashes != {name: digest(output / name) for name in SOURCES}:
            raise AssertionError('CSV changed during report generation')
        result['artifact_sha256'] = {name: digest(temporary / name) for name in paths if name != 'report_verification.json'}
        (temporary / 'report_verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        (output / 'figures').mkdir(exist_ok=True)
        for name in paths:
            os.replace(temporary / name, output / name)
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', required=True)
    args = parser.parse_args()
    result = render(args.output_dir)
    print(json.dumps({'status': result['status'], 'tables_checked': result['word']['tables_checked'],
                      'table_cells_checked': result['word']['table_cells_checked']}, ensure_ascii=False))
