"""Render the fixed local-baseline MELD comparison, and preserve earlier reports.

No inference is performed. Source CSVs are immutable during rendering. Each Word
table is independently checked against CSV cells; the combined document retains
the complete original body and remaps appended image relationships explicitly.
"""
import argparse
from copy import deepcopy
from io import BytesIO
import json
import math
import os
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET
import zipfile

import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from lxml import etree

try:
    from .render_official_reliability import _set_run, three_line_table, read_csv, digest, FONT, NS
    from .render_public_m3net_comparison import display, verify_word
except ImportError:
    from render_official_reliability import _set_run, three_line_table, read_csv, digest, FONT, NS
    from render_public_m3net_comparison import display, verify_word

DOCX_NAME = 'CMERC_MELD_实验二_分类纠错与校准对比.docx'
COMBINED_NAME = 'CMERC_官方Checkpoint_实验一至六综合报告_含MELD实验二.docx'
FIGURE_STEM = 'figures/reliability_MELD_baseline_comparison'
SOURCES = ('paired_predictions.csv', 'classification_transitions.csv',
    'high_confidence_transitions.csv', 'auxiliary_threshold_transitions.csv',
    'classwise_transitions.csv', 'calibration_comparison.csv',
    'calibration_bins.csv', 'typical_cases.csv')
EMOTIONS = ('neutral', 'surprise', 'fear', 'sadness', 'joy', 'disgust', 'anger')
REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def make_figure(directory, rows):
    models = list(dict.fromkeys(row['model'] for row in rows))
    if len(models) != 2:
        raise ValueError('Expected two models')
    records = []
    with plt.rc_context({'font.family': 'DejaVu Sans', 'font.size': 9,
                         'svg.fonttype': 'none', 'pdf.fonttype': 42,
                         'axes.spines.top': False, 'axes.spines.right': False}):
        fig, (axis, support) = plt.subplots(2, 1, figsize=(6.2, 6.8),
            gridspec_kw={'height_ratios': [3.4, 1]})
        fig.subplots_adjust(left=.15, right=.97, top=.94, bottom=.09, hspace=.40)
        axis.plot([0, 1], [0, 1], '--', color='#777777', linewidth=1, label='Ideal calibration')
        for index, (model, color, marker) in enumerate(zip(models, ('#326A98', '#AF6B38'), ('o', 's'))):
            bins = sorted((r for r in rows if r['model'] == model), key=lambda r: int(r['bin_index']))
            if [int(r['bin_index']) for r in bins] != list(range(10)):
                raise ValueError('Missing/duplicate calibration bins')
            if sum(int(r['count']) for r in bins) != 2610:
                raise ValueError('Bin support must total 2610')
            active = [r for r in bins if int(r['count']) > 0]
            points = [[float(r['mean_confidence']), float(r['accuracy'])] for r in active]
            if any(not all(math.isfinite(v) and 0 <= v <= 1 for v in p) for p in points):
                raise ValueError('Invalid reliability point')
            axis.plot([p[0] for p in points], [p[1] for p in points], color=color,
                      marker=marker, markersize=4, linewidth=1.2, label=model)
            support.bar([(float(r['lower_bound']) + float(r['upper_bound'])) / 2
                         + (-.021 if index == 0 else .021) for r in bins],
                        [int(r['count']) for r in bins], width=.038, color=color)
            records.append(dict(model=model, counts=[int(r['count']) for r in bins],
                nonempty_bin_indices=[int(r['bin_index']) for r in active], points=points))
        axis.set(xlim=(0, 1), ylim=(0, 1), xlabel='Mean confidence',
                 ylabel='Empirical accuracy', title='MELD: fixed checkpoints, descriptive comparison')
        axis.set_aspect('equal', adjustable='box')
        axis.legend(loc='upper left', frameon=False, fontsize=7)
        support.set(xlim=(0, 1), xlabel='Confidence bin', ylabel='Utterances')
        for ext in ('png', 'svg', 'pdf'):
            fig.savefig(directory / f'{FIGURE_STEM}.{ext}', dpi=300, facecolor='white')
        plt.close(fig)
    return dict(source_csv='calibration_bins.csv', ideal_diagonal=True,
                empty_bins_omitted=True, model_series=records)


def make_word(directory, sources, audit_text):
    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Mm(297), Mm(210)
        section.top_margin = section.bottom_margin = Mm(18)
        section.left_margin = section.right_margin = Mm(18)
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Caption'):
        style = document.styles[name]
        style.font.name, style.font.size = FONT, Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn(f'w:{slot}'), FONT)
    tables, headings = [], []

    def para(text, heading=False):
        paragraph = document.add_paragraph()
        _set_run(paragraph.add_run(text), bold=heading)
        if heading:
            headings.append(text)

    def table(title, source, columns, headers, *, percents=(), integers=(), signed=(), indices=None):
        fields, rows = sources[source]
        if set(columns) - set(fields):
            raise ValueError(f'Missing report columns: {source}: {set(columns)-set(fields)}')
        chosen = list(range(len(rows))) if indices is None else indices
        numeric = set(percents) | set(integers) | set(signed)
        rendered = [[display(rows[i][field], percent=field in percents,
            integer=field in integers, signed=field in signed)
            if field in numeric or field in {'brier_score', 'nll'}
            or field.startswith(('baseline_probability_', 'cmerc_probability_')) else rows[i][field]
            for field in columns] for i in chosen]
        para(title, True)
        tables.append(dict(table_index=len(document.tables), source_csv=source,
            columns=columns, source_row_indices=chosen, headers=headers,
            percentage_columns=list(percents), integer_columns=list(integers),
            signed_columns=list(signed), rendered_rows=rendered))
        three_line_table(document, headers, rendered)

    para('CMERC MELD 实验二：分类纠错与概率校准对比', True)
    para('本报告将预先固定的本地 README 复现 Baseline（历史测试集选模）与作者官方 CMERC '
         '进行相同测试协议下的描述性行为比较。MELD 七分类测试集包含 2610 条有效话语；'
         'CMERC 复用已验证预测，Baseline 为本轮实际完整三模态推理；不训练、不换 checkpoint、'
         '不对概率进行温度缩放或重新归一化。')
    para('一、数据来源、历史选模风险与可比性', True)
    para('Baseline 为本地 experiments/official_readme_reproduction/checkpoints/'
         'MELD_baseline_seed67137/best.pt，训练 seed=67137，选中 epoch=5，'
         '采用 official_test_selection。它不是 M³Net 作者公开 MELD 权重，'
         '不能描述为独立验证集选模模型。历史测试集选模使测试成绩存在选择偏倚；'
         '本轮固定该候选，不比较其他候选的测试 F1 进行替换。')
    para('CMERC 为作者公开 MELD_checkpoint.pkl，完整输入性能 W-F1=66.8484%。'
         '模型结构、特征、测试顺序、Batch=16、Padding/Mask、七分类映射、cuDNN RNN 与'
         '两处既有 PyG 兼容补丁按运行清单核验。训练配对、作者 seed/epoch/选模历史仍未知。'
         '故可比较这两份固定权重的预测行为，不能归因 CL、HSCL、CC 的因果收益或称严格模块消融。')
    para('审计原文（独立保存的 MELD_baseline_checkpoint_audit.md）', True)
    for line in audit_text.splitlines():
        if line.strip():
            para(line)
    para('二、逐样本分类变化', True)
    table('表 1　2610 条话语的互斥分类转移', 'classification_transitions.csv',
          ['transition', 'count', 'ratio'], ['分类转移', '数量', '占全部话语 (%)'],
          percents=('ratio',), integers=('count',))
    base, cm = sources['calibration_comparison.csv'][1]
    para(f'Baseline 错误数 {base["total_errors"]}；官方 CMERC 错误数 {cm["total_errors"]}。'
         f'CMERC−Baseline 的 W-F1 差值为 '
         f'{display(float(cm["weighted_f1"])-float(base["weighted_f1"]), percent=True, signed=True)} 个百分点；'
         f'净错误数变化为 {int(cm["total_errors"])-int(base["total_errors"]):+d}。'
         '纠正与退化同时报告，置信度下降不等于分类纠正。')
    para('三、高置信度错误与固定辅助阈值', True)
    for source, title in (('high_confidence_transitions.csv', '原实验阈值 ≥0.80、≥0.90、≥0.95'),
                          ('auxiliary_threshold_transitions.csv', '预先固定辅助阈值 ≥0.40、≥0.50、≥0.60、≥0.70')):
        table(f'表　{title}：预测支持数和错误率', source,
              ['threshold', 'baseline_high_confidence_predictions', 'baseline_high_confidence_errors',
               'baseline_high_confidence_error_rate', 'cmerc_high_confidence_predictions',
               'cmerc_high_confidence_errors', 'cmerc_high_confidence_error_rate'],
              ['阈值', 'Baseline 预测数', 'Baseline 错误数', 'Baseline 错误率 (%)',
               'CMERC 预测数', 'CMERC 错误数', 'CMERC 错误率 (%)'],
              percents=('baseline_high_confidence_error_rate', 'cmerc_high_confidence_error_rate'),
              integers=('baseline_high_confidence_predictions', 'baseline_high_confidence_errors',
                        'cmerc_high_confidence_predictions', 'cmerc_high_confidence_errors'))
        table(f'表　{title}：原错误去向与新增错误', source,
              ['threshold', 'baseline_high_confidence_errors', 'corrected', 'still_wrong_confidence_lower',
               'still_wrong_confidence_not_lower', 'still_wrong_above_threshold', 'new_cmerc_hc_errors'],
              ['阈值', '原高置信错误', '分类纠正', '仍错置信度下降', '仍错置信度未降',
               '仍错且过原阈值', 'CMERC 新增高置信错误'],
              integers=('baseline_high_confidence_errors', 'corrected', 'still_wrong_confidence_lower',
                        'still_wrong_confidence_not_lower', 'still_wrong_above_threshold', 'new_cmerc_hc_errors'))
        table(f'表　{title}：原错误去向比例', source,
              ['threshold', 'corrected_ratio', 'still_wrong_confidence_lower_ratio',
               'still_wrong_confidence_not_lower_ratio', 'still_wrong_above_threshold_ratio'],
              ['阈值', '纠正 (%)', '仍错降置信 (%)', '仍错未降 (%)', '仍错过阈值 (%)'],
              percents=('corrected_ratio', 'still_wrong_confidence_lower_ratio',
                        'still_wrong_confidence_not_lower_ratio', 'still_wrong_above_threshold_ratio'))
        table(f'表　{title}：新增错误的分母和来源', source,
              ['threshold', 'eligible_non_baseline_hc_errors', 'new_cmerc_hc_errors',
               'new_cmerc_hc_error_ratio_of_eligible', 'new_cmerc_hc_errors_from_baseline_correct',
               'new_cmerc_hc_errors_from_baseline_wrong_below_threshold'],
              ['阈值', '原非高置信错误数', '新增数量', '新增占可进入样本 (%)',
               '来自原正确', '来自原低置信错误'],
              percents=('new_cmerc_hc_error_ratio_of_eligible',),
              integers=('eligible_non_baseline_hc_errors', 'new_cmerc_hc_errors',
                        'new_cmerc_hc_errors_from_baseline_correct',
                        'new_cmerc_hc_errors_from_baseline_wrong_below_threshold'))
    para('表注：每个阈值均采用 ≥。错误率分母为该模型达到阈值的全部预测数，'
         '分母为零写 N/A；零高置信错误不能解释成可靠性完美。纠正/仍错降置信/仍错未降三类互斥，'
         '比例分母为原 Baseline 高置信错误数；“仍错且过原阈值”为重叠子集。新增比例分母为'
         '2610−原 Baseline 高置信错误数，新增来源两项之和等于新增数量。辅助阈值独立展示，'
         '不替代原实验阈值，也不根据结果选择最有利阈值。')
    para('最大预测置信度与真实类别概率在配对 CSV 分字段保留。错误预测类别可能变化，'
         '最大概率下降不等于同一错误类别概率下降，不据此认定模态冲突。')
    table('表　七类真实情绪的纠正与退化', 'classwise_transitions.csv',
          ['true_emotion', 'support', 'baseline_errors', 'cmerc_errors', 'corrected', 'regressed',
           'both_wrong', 'net_error_change'],
          ['真实情绪', '支持数', 'Baseline 错误', 'CMERC 错误', '错→对', '对→错', '都错', '净错误变化'],
          integers=('support', 'baseline_errors', 'cmerc_errors', 'corrected', 'regressed',
                    'both_wrong', 'net_error_change'))
    para('表注：净错误变化=CMERC 错误−Baseline 错误；正值表示错误增加。'
         'fear、sadness、disgust 与其余类别一并据实展示，不预设改善方向。')
    para('四、完整测试集的概率校准对比', True)
    table('表　分类性能与平均置信度', 'calibration_comparison.csv',
          ['model', 'valid_utterances', 'total_errors', 'accuracy', 'weighted_f1',
           'mean_confidence', 'mean_error_confidence'],
          ['模型', '话语数', '错误数', 'Accuracy (%)', 'W-F1 (%)', '平均置信度 (%)', '错误平均置信度 (%)'],
          percents=('accuracy', 'weighted_f1', 'mean_confidence', 'mean_error_confidence'),
          integers=('valid_utterances', 'total_errors'))
    table('表　概率校准指标', 'calibration_comparison.csv',
          ['model', 'ece', 'mce', 'brier_score', 'nll'],
          ['模型', 'ECE (%)', 'MCE (%)', 'Brier Score', 'NLL'], percents=('ece', 'mce'))
    para('表注：全部 2610 条有效话语参与计算。ECE 为十等宽箱 [0,0.1)、…、[0.9,1] 的'
         '样本占比乘绝对校准差值之和；MCE 为非空箱绝对差值最大值；空箱不构造观测。'
         'Brier 是逐样本所有类别平方误差之和再平均，不除以类别数或 2。'
         'NLL 是真实类别概率负自然对数的均值，不裁剪、不重新归一化概率。'
         'Signed Gap=平均置信度−准确率，正为过度自信，负为置信度不足。')
    for model in dict.fromkeys(r['model'] for r in sources['calibration_bins.csv'][1]):
        selected = [i for i, r in enumerate(sources['calibration_bins.csv'][1]) if r['model'] == model]
        table(f'表　{model} 十等宽校准分箱', 'calibration_bins.csv',
              ['lower_bound', 'upper_bound', 'count', 'mean_confidence', 'accuracy', 'signed_gap'],
              ['下界', '上界', '支持数', '平均置信度 (%)', '实际准确率 (%)', '差值 (pp)'],
              percents=('mean_confidence', 'accuracy', 'signed_gap'), integers=('count',),
              signed=('signed_gap',), indices=selected)
    document.add_picture(str(directory / f'{FIGURE_STEM}.png'), width=Inches(5.6))
    para('图　MELD 固定权重可靠性对比。虚线为理想对角线；仅非空箱绘制真实观测点，'
         '连线仅引导阅读，下方显示全部十箱支持数。图不拟合校准器，未提供跨种子不确定性。')
    para('五、典型案例：按固定规则展示正反两类变化', True)
    cases = sources['typical_cases.csv'][1]
    case_types = ('corrected_baseline_hc_error', 'still_wrong_confidence_lower',
                  'baseline_correct_new_cmerc_hc_error', 'both_high_confidence_wrong')
    if {r['case_type'] for r in cases} - set(case_types):
        raise ValueError('Unknown case category')
    para('所有符合条件案例均保存在 typical_cases.csv（含 threshold_group），正文固定使用辅助'
         '阈值 0.60，每种类型按 selection_score 降序及稳定 ID 排序展示前两条。'
         'selection_score：纠正及仍错降置信使用 Baseline 最大概率；新增错误使用 CMERC 最大概率；'
         '双方高置信错误使用两者最大概率较小值。原实验 0.90 阈值无案例的类型明确报告无，'
         '不因稀少而改阈值；展示排序不参与选模。')
    for case_type in case_types:
        primary = [r for r in cases if r['case_type'] == case_type and float(r['threshold']) == .9]
        para(f'案例类型 {case_type}：原阈值 0.90 匹配 {len(primary)} 条。', True)
        selected = sorted((i for i, r in enumerate(cases)
                           if r['case_type'] == case_type and float(r['threshold']) == .6),
                          key=lambda i: (-float(cases[i]['selection_score']), cases[i]['dialogue_id'],
                                         int(cases[i]['utterance_index'])))[:2]
        if not selected:
            para('辅助阈值 0.60 下无符合条件案例。')
        for index in selected:
            row = cases[index]
            para(f'对话 {row["dialogue_id"]}，索引 {row["utterance_index"]}，源 ID '
                 f'{row["source_utterance_id"]}，真实情绪 {row["true_emotion"]}；原文：{row["utterance_text"]}')
            table('案例预测与两种概率', 'typical_cases.csv',
                  ['baseline_predicted_emotion', 'baseline_predicted_confidence', 'baseline_true_label_probability',
                   'cmerc_predicted_emotion', 'cmerc_predicted_confidence', 'cmerc_true_label_probability'],
                  ['Baseline 预测', 'Baseline 最大概率 (%)', 'Baseline 真实类概率 (%)',
                   'CMERC 预测', 'CMERC 最大概率 (%)', 'CMERC 真实类概率 (%)'],
                  percents=('baseline_predicted_confidence', 'baseline_true_label_probability',
                            'cmerc_predicted_confidence', 'cmerc_true_label_probability'), indices=[index])
            for prefix, label in (('baseline', 'Baseline'), ('cmerc', 'CMERC')):
                table(f'{label} 全部七类概率', 'typical_cases.csv',
                      [f'{prefix}_probability_{emotion}' for emotion in EMOTIONS], list(EMOTIONS), indices=[index])
    para('六、结论与限制', True)
    transitions = {r['transition']: int(r['count']) for r in sources['classification_transitions.csv'][1]}
    para(f'分类纠正 {transitions.get("baseline_wrong_cmerc_correct", 0)} 条，新增分类错误 '
         f'{transitions.get("baseline_correct_cmerc_wrong", 0)} 条；两者都正确 '
         f'{transitions.get("both_correct", 0)} 条，都错误 {transitions.get("both_wrong", 0)} 条。'
         '分类上的小幅净变化与置信度上的变化是不同问题。')
    para('CMERC−Baseline 的 ECE 差值 '
         f'{display(float(cm["ece"])-float(base["ece"]), percent=True, signed=True)} 个百分点；MCE '
         f'{display(float(cm["mce"])-float(base["mce"]), percent=True, signed=True)} 个百分点；Brier '
         f'{display(float(cm["brier_score"])-float(base["brier_score"]), signed=True)}；NLL '
         f'{display(float(cm["nll"])-float(base["nll"]), signed=True)}。'
         '应结合分类转移、所有阈值及完整校准评分判断，不能凭没有 ≥0.90 错误宣称整体可靠。')
    direction = '上升' if float(cm['mean_error_confidence']) > float(base['mean_error_confidence']) else '下降或不变'
    para(f'错误样本平均置信度从 {display(base["mean_error_confidence"], percent=True)}% 到 '
         f'{display(cm["mean_error_confidence"], percent=True)}%，方向为{direction}。'
         '因此整体概率评分的变化不能直接解释成模型普遍压低了错误预测的置信度。')
    auxiliary = {float(r['threshold']): r for r in sources['auxiliary_threshold_transitions.csv'][1]}
    if .6 in auxiliary:
        row = auxiliary[.6]
        para('预先固定辅助阈值 0.60 下，Baseline 达阈值预测 '
             f'{row["baseline_high_confidence_predictions"]} 条、其中错误 '
             f'{row["baseline_high_confidence_errors"]} 条，错误率 '
             f'{display(row["baseline_high_confidence_error_rate"], percent=True)}%；CMERC 分别为 '
             f'{row["cmerc_high_confidence_predictions"]} 条、'
             f'{row["cmerc_high_confidence_errors"]} 条、'
             f'{display(row["cmerc_high_confidence_error_rate"], percent=True)}%。'
             '覆盖样本数和条件错误率必须一并解释，不能声称全部可靠性指标都得到改善。')
    para('本报告仅支持固定权重的描述性测试行为对照。本地 Baseline 历史测试集选模风险、'
         '与官方 CMERC 未确认的训练配对关系均保留。IEMOCAP 实验二使用作者公开 IEMOCAP3 权重，'
         'MELD 使用本地 README 复现权重，两者来源不同，不合并为同一组严格配对模块消融。')
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            _set_run(run)
    document.save(directory / DOCX_NAME)
    return tables, headings


def _body_elements(document):
    return [element for element in document.element.body if element.tag != qn('w:sectPr')]


def _canonical(element):
    return etree.tostring(element, method='c14n')


def _image_sequence(document):
    result = []
    for element in document.element.body.iter():
        relationship = element.get(f'{{{REL_NS}}}embed')
        if relationship:
            part = document.part.related_parts[relationship]
            import hashlib
            result.append(hashlib.sha256(part.blob).hexdigest())
    return result


def append_document(target, source):
    """Append whole body elements, retaining images and external hyperlinks."""
    image_mapping = {}
    max_drawing = max((int(e.get('id')) for e in target.element.body.iter()
                       if e.tag.endswith('}docPr') and e.get('id', '').isdigit()), default=0)
    count = 0
    for source_element in _body_elements(source):
        copied = deepcopy(source_element)
        for element in copied.iter():
            for attribute, old_id in list(element.attrib.items()):
                if not attribute.startswith(f'{{{REL_NS}}}'):
                    continue
                if old_id not in image_mapping:
                    relationship = source.part.rels[old_id]
                    if relationship.reltype.endswith('/image'):
                        new_id, _ = target.part.get_or_add_image(BytesIO(relationship.target_part.blob))
                    elif relationship.is_external:
                        new_id = target.part.relate_to(relationship.target_ref,
                                                      relationship.reltype, is_external=True)
                    else:
                        raise ValueError(f'Unsupported internal relation: {relationship.reltype}')
                    image_mapping[old_id] = new_id
                element.set(attribute, image_mapping[old_id])
            if element.tag.endswith('}docPr'):
                max_drawing += 1
                element.set('id', str(max_drawing))
        target.element.body.insert(len(target.element.body) - 1, copied)
        count += 1
    return count


def _text_and_tables(document):
    return [''.join(e.itertext()) for e in _body_elements(document)]


def make_combined(base_path, iem_path, meld_path, destination):
    paths = [Path(base_path), Path(iem_path), Path(meld_path)]
    original_hashes = {str(p): digest(p) for p in paths}
    base, iem, meld = [Document(p) for p in paths]
    prefix = [_canonical(e) for e in _body_elements(base)]
    original_images = _image_sequence(base)
    original_tables = [[[c.text for c in row.cells] for row in t.rows] for t in base.tables]
    p = base.add_paragraph()
    _set_run(p.add_run('后续补充导读：实验二的对照来源与历史状态'), bold=True)
    p = base.add_paragraph()
    _set_run(p.add_run('前文完整保留原正式报告及其当时 Baseline pending 的历史说明。'
        '本次在后追加已完成的 IEMOCAP 与 MELD 实验二，替代其“待补充”的当前状态。'
        'IEMOCAP 使用 M³Net 作者公开 Roberta_IEMOCAP3.pth.tar；MELD 使用本地 README '
        '复现、epoch 5、official_test_selection 的 Baseline。两数据集对照权重来源不同，'
        '均不构成严格训练配对，不能归因为 CL、HSCL、CC 的模块因果收益。'))
    intro_count = len(_body_elements(base)) - len(prefix)
    iem_start = len(_body_elements(base))
    iem_count = append_document(base, iem)
    meld_start = len(_body_elements(base))
    meld_count = append_document(base, meld)
    base.save(destination)
    combined = Document(destination)
    elements = _body_elements(combined)
    if [_canonical(e) for e in elements[:len(prefix)]] != prefix:
        raise AssertionError('Original report body prefix changed')
    if [[[c.text for c in row.cells] for row in t.rows] for t in combined.tables[:len(original_tables)]] != original_tables:
        raise AssertionError('Original report table values changed')
    expected_images = original_images + _image_sequence(iem) + _image_sequence(meld)
    if _image_sequence(combined) != expected_images:
        raise AssertionError('Combined report image relationship or bytes lost')
    for source, start, count in ((iem, iem_start, iem_count), (meld, meld_start, meld_count)):
        actual = [''.join(e.itertext()) for e in elements[start:start+count]]
        if actual != _text_and_tables(source):
            raise AssertionError('Appended report text/table content changed')
    if original_hashes != {str(p): digest(p) for p in paths}:
        raise AssertionError('A source report changed')
    with zipfile.ZipFile(destination) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
        if any(name.startswith(('word/header', 'word/footer')) for name in archive.namelist()):
            raise AssertionError('Unexpected combined header/footer parts')
        for run in root.findall('.//w:r', NS):
            fonts, color, size = (run.find('w:rPr/'+name, NS) for name in ('w:rFonts', 'w:color', 'w:sz'))
            if (fonts is None or any(fonts.get(qn('w:'+slot)) != FONT for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'))
                    or color is None or color.get(qn('w:val')) != '000000'
                    or size is None or size.get(qn('w:val')) != '24'):
                raise AssertionError('Combined Song/12pt/black style mismatch')
        for table in root.findall('.//w:tbl', NS):
            borders = table.find('w:tblPr/w:tblBorders', NS)
            if borders is None:
                raise AssertionError('Missing combined three-line table border')
            for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
                edge = borders.find('w:'+side, NS)
                if edge is None or edge.get(qn('w:val')) != ('single' if side in ('top', 'bottom') else 'nil'):
                    raise AssertionError('Combined table border changed')
            for cell in table.findall('w:tr', NS)[0].findall('w:tc', NS):
                edge = cell.find('w:tcPr/w:tcBorders/w:bottom', NS)
                if edge is None or edge.get(qn('w:val')) != 'single':
                    raise AssertionError('Combined table header rule missing')
    return dict(status='passed', source_sha256=original_hashes, source_reports_unchanged=True,
        original_body_prefix_identical=True, original_body_elements=len(prefix),
        original_tables_preserved=len(original_tables), original_images_preserved=len(original_images),
        appended_iem_body_elements=iem_count, appended_meld_body_elements=meld_count,
        added_intro_elements=intro_count, all_appended_text_and_tables_identical=True,
        all_embedded_image_sha256=expected_images, font=FONT, font_size_points=12,
        color='000000', explicit_three_line_tables=True, headers_footers=False, office_layout_rendered=False,
        layout_note='OOXML values/styles/relationships checked; Office pagination not rendered.',
        combined_sha256=digest(destination))


def render(output):
    output = Path(output).resolve()
    sources = {name: read_csv(output / name) for name in SOURCES}
    hashes = {name: digest(output / name) for name in SOURCES}
    audit_path = output / 'MELD_baseline_checkpoint_audit.md'
    hashes[audit_path.name] = digest(audit_path)
    base_path = output.parent / 'CMERC_官方Checkpoint_可靠性诊断报告.docx'
    iem_path = output.parent / 'experiment_2_public_m3net_comparison/CMERC_官方权重_实验二分类纠错与校准对比.docx'
    paths = [DOCX_NAME, COMBINED_NAME, 'report_verification.json', 'combined_report_verification.json',
             'figure_contract.md'] + [f'{FIGURE_STEM}.{ext}' for ext in ('png', 'svg', 'pdf')]
    for name in paths:
        if (output / name).exists():
            raise FileExistsError(f'Refusing to overwrite report artifact: {output / name}')
    with tempfile.TemporaryDirectory(prefix='.meld_render_', dir=output) as folder:
        temporary = Path(folder)
        (temporary / 'figures').mkdir()
        (temporary / 'figure_contract.md').write_text(
            '# MELD comparison figure contract\n\nPython/matplotlib; ten equal-width bins, '
            '2610 valid utterances per model. Local README Baseline (test-selected) versus '
            'Official CMERC. Mean maximum probability against empirical accuracy, ideal '
            'diagonal, no empty-bin point; support bars below. Lines guide the eye only. '
            'No calibration fitting, causal module claim, or seed-level uncertainty. '
            'PNG 300dpi, editable-text SVG, TrueType PDF.\n', encoding='utf-8')
        figure = make_figure(temporary, sources['calibration_bins.csv'][1])
        records, headings = make_word(temporary, sources, audit_path.read_text(encoding='utf-8'))
        verification = dict(status='passed', training=False, gpu_inference=False,
            source_csv_sha256=hashes, figure=figure,
            word=verify_word(temporary / DOCX_NAME, sources, records, headings,
                             temporary / f'{FIGURE_STEM}.png'))
        combined = make_combined(base_path, iem_path, temporary / DOCX_NAME, temporary / COMBINED_NAME)
        combined['source_sha256'][str(output / DOCX_NAME)] = combined['source_sha256'].pop(str(temporary / DOCX_NAME))
        combined['appendix_csv_verification'] = {'IEMOCAP': str(iem_path.parent / 'report_verification.json'),
            'MELD': 'report_verification.json', 'MELD_tables_checked': verification['word']['tables_checked'],
            'MELD_table_cells_checked': verification['word']['table_cells_checked']}
        previous = json.loads((output.parent / 'report_verification.json').read_text())
        if previous['status'] != 'passed' or not previous['word']['all_table_cells_match_csv']:
            raise AssertionError('Previous official diagnostics report not verified')
        if previous['outputs'][base_path.name]['sha256'] != digest(base_path):
            raise AssertionError('Previous official diagnostics Word changed since verification')
        for name, sha in previous['sources'].items():
            if digest(output.parent / name) != sha:
                raise AssertionError('Previous official diagnostics CSV changed since verification')
        combined['original_csv_provenance_unchanged'] = True
        for dataset, folder_name in (('IEMOCAP', iem_path.parent), ('MELD', output)):
            if dataset == 'IEMOCAP':
                result = json.loads((folder_name / 'report_verification.json').read_text())
                if result['status'] != 'passed' or not result['word']['all_table_cells_match_csv']:
                    raise AssertionError('Prior IEMOCAP report not verified')
                if result['artifact_sha256'][iem_path.name] != digest(iem_path):
                    raise AssertionError('Prior IEMOCAP Word changed since verification')
                for name, sha in result['source_csv_sha256'].items():
                    if digest(folder_name / name) != sha:
                        raise AssertionError('Prior IEMOCAP CSV changed since report verification')
        if hashes != {name: digest(output / name) for name in hashes}:
            raise AssertionError('Input CSV or audit changed during rendering')
        (temporary / 'combined_report_verification.json').write_text(json.dumps(combined, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        verification['artifact_sha256'] = {name: digest(temporary / name) for name in paths if name != 'report_verification.json'}
        (temporary / 'report_verification.json').write_text(json.dumps(verification, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        (output / 'figures').mkdir(exist_ok=True)
        for name in paths:
            os.replace(temporary / name, output / name)
    return verification


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', required=True)
    args = parser.parse_args()
    result = render(args.output_dir)
    print(json.dumps({'status': result['status'], 'tables_checked': result['word']['tables_checked'],
                      'table_cells_checked': result['word']['table_cells_checked']}, ensure_ascii=False))
