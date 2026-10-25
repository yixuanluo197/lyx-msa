"""Render the fixed author-checkpoint diagnostics without inference or training.

All statistics come from the independent official_checkpoint_analysis CSVs.
The Python-only figures show observed nonempty calibration bins, not fitted
calibration curves. The generated workbook is verified directly from OOXML so
no spreadsheet reader dependency is needed.
"""
import argparse
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET
import zipfile

import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
import xlsxwriter

DATASETS = ('IEMOCAP', 'MELD')
DOCX_NAME = 'CMERC_官方Checkpoint_可靠性诊断报告.docx'
XLSX_NAME = 'CMERC_官方Checkpoint_诊断统计.xlsx'
FONT = '宋体'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
NUMERIC_FIELDS = {
    'utterance_index', 'true_label', 'predicted_label', 'correct', 'threshold',
    'support', 'total_errors', 'error_rate', 'valid_utterances', 'accuracy',
    'weighted_f1', 'macro_f1', 'mean_confidence', 'mean_error_confidence',
    'median_error_confidence', 'ece', 'mce', 'brier_score', 'nll', 'bin_index',
    'lower_bound', 'upper_bound', 'count', 'signed_gap', 'absolute_gap',
    'weight', 'ece_contribution', 'error_count', 'ratio_of_all_errors',
    'predicted_confidence', 'true_label_probability', 'true_label_log_probability',
    'full_correct_samples', 'full_correct_count', 'ratio', 'full_prediction',
    'full_predicted_label', 'full_confidence', 'masked_prediction',
    'masked_predicted_label', 'masked_confidence', 'masked_correct',
    'masked_true_label_probability', 'full_correct_to_masked_high_confidence_error_count',
    'full_correct_to_masked_error_count', 'full_wrong_still_wrong_count',
    'new_high_confidence_errors', 'full_correct_to_masked_high_confidence_errors',
    'full_wrong_still_wrong', 'valid_samples', 'full_wrong_samples',
    'full_correct', 'full_true_label_probability', 'true_label_probability_delta',
    'still_wrong_count', 'still_wrong_ratio',
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        fields, rows = reader.fieldnames, list(reader)
    if not fields or len(fields) != len(set(fields)):
        raise ValueError(f'Invalid CSV header: {path}')
    if any(None in row or any(value is None for value in row.values()) for row in rows):
        raise ValueError(f'Malformed CSV row: {path}')
    return fields, rows


def numeric_field(name):
    return name in NUMERIC_FIELDS or name.startswith((
        'probability_', 'errors_ge_', 'high_confidence_predictions_ge_',
        'high_confidence_error_rate_ge_'))


def excel_value(field, value):
    if value in ('', 'None', 'N/A', 'null'):
        return value
    if numeric_field(field):
        if value in ('True', 'False'):
            return int(value == 'True')
        number = float(value)
        if not math.isfinite(number):
            return value
        if number.is_integer() and abs(number) < 2**53:
            return int(number)
        return number
    return value


def formatted(value, percent=False, signed=False):
    if value in (None, '', 'None', 'N/A', 'null'):
        return 'N/A'
    number = float(value) * (100 if percent else 1)
    if math.isinf(number):
        return '+∞' if number > 0 else '−∞'
    if number.is_integer() and not percent and not signed:
        return str(int(number))
    return format(number, '+.4f' if signed else '.4f' if percent else '.6f')


def _set_run(run, *, bold=None):
    run.font.name = FONT
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        properties.insert(0, fonts)
    for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{script}'), FONT)


def three_line_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'single' if side in ('top', 'bottom') else 'nil')
        edge.set(qn('w:color'), '000000')
        edge.set(qn('w:sz'), '8')
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
        edges, edge = OxmlElement('w:tcBorders'), OxmlElement('w:bottom')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:color'), '000000')
        edge.set(qn('w:sz'), '4')
        edges.append(edge)
        cell._tc.get_or_add_tcPr().append(edges)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = str(text)
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run(run, bold=index == 0)
    return table


def figure_contract(groups):
    return ('# Official checkpoint reliability figure contract\n\n'
            'Backend: Python / matplotlib exclusively. Archetype: quantitative grid.\n\n'
            'Claim: describe the confidence–accuracy gaps of the two fixed author '
            'checkpoints on their complete test sets; no Baseline improvement claim.\n\n'
            'Evidence: nonempty-bin mean confidence versus accuracy on equal [0,1] '
            'axes, ideal calibration diagonal, plus all ten bin supports below. '
            'Empty bins have no reliability point. Lines only guide the eye.\n\n'
            'Exports: 4.6 × 5.6 inches; PNG 300 dpi for Word, editable-text SVG and '
            'TrueType PDF. Source data are each dataset/calibration_bins.csv.\n\n'
            'Risks: sparse bins are unstable; no fitted calibration, confidence '
            'intervals, cross-seed claims or causal modality-conflict claims.\n\n' +
            '\n'.join(f'- {dataset}: n={sum(int(r["count"]) for r in rows)}; '
                      f'nonempty bins={sum(int(r["count"]) > 0 for r in rows)}'
                      for dataset, rows in groups.items()) + '\n')


def make_figures(directory, groups):
    records = []
    with plt.rc_context({'font.family': 'DejaVu Sans', 'font.size': 8,
                         'svg.fonttype': 'none', 'pdf.fonttype': 42,
                         'axes.spines.top': False, 'axes.spines.right': False}):
        for dataset, bins in groups.items():
            active = [row for row in bins if int(row['count']) > 0]
            fig, (axis, support) = plt.subplots(2, 1, figsize=(4.6, 5.6),
                gridspec_kw={'height_ratios': [3, 1]})
            fig.subplots_adjust(left=.18, right=.96, top=.92, bottom=.10, hspace=.45)
            axis.plot([0, 1], [0, 1], color='#777777', linestyle='--', linewidth=.9,
                      label='Ideal calibration')
            x = [float(row['mean_confidence']) for row in active]
            y = [float(row['accuracy']) for row in active]
            axis.plot(x, y, color='#346B95', marker='o', markersize=4, linewidth=1.1,
                      label='Official CMERC')
            axis.set(xlim=(0, 1), ylim=(0, 1), xlabel='Mean confidence',
                     ylabel='Empirical accuracy', title=dataset)
            axis.set_aspect('equal', adjustable='box')
            axis.set_xticks([0, .2, .4, .6, .8, 1])
            axis.set_yticks([0, .2, .4, .6, .8, 1])
            axis.legend(loc='upper left', frameon=False, fontsize=7)
            support.bar([(float(r['lower_bound']) + float(r['upper_bound'])) / 2 for r in bins],
                        [int(r['count']) for r in bins], width=.085, color='#346B95', alpha=.8)
            support.set(xlim=(0, 1), xlabel='Confidence bin', ylabel='Count')
            support.set_xticks([0, .2, .4, .6, .8, 1])
            paths = []
            for extension in ('png', 'svg', 'pdf'):
                name = f'figures/reliability_{dataset}.{extension}'
                fig.savefig(directory / name, dpi=300, facecolor='white')
                paths.append(name)
            plt.close(fig)
            records.append(dict(dataset=dataset, files=paths,
                source_csv=f'{dataset}/calibration_bins.csv',
                bin_indices=[int(r['bin_index']) for r in active],
                points=list(map(list, zip(x, y))), counts=[int(r['count']) for r in bins],
                ideal_diagonal=True, empty_bins_omitted=True))
    return records


def make_word(directory, sources, baseline_text):
    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Mm(297), Mm(210)
        section.top_margin = section.bottom_margin = Mm(18)
        section.left_margin = section.right_margin = Mm(18)
    for style in document.styles:
        if style.type == 1:
            style.font.name = FONT
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
    table_records = []

    def paragraph(text, heading=False):
        paragraph = document.add_paragraph()
        _set_run(paragraph.add_run(text), bold=heading)
        return paragraph

    def table(title, source, columns, labels=None, percents=(), signed=(), selected=None):
        fields, source_rows = sources[source]
        if any(field not in fields for field in columns):
            raise ValueError(f'Missing Word columns in {source}: {set(columns) - set(fields)}')
        indices = list(range(len(source_rows))) if selected is None else selected
        rows = [[formatted(source_rows[i][field], field in percents, field in signed)
                 if numeric_field(field) else source_rows[i][field] for field in columns]
                for i in indices]
        paragraph(title, True)
        index = len(document.tables)
        headers = labels or columns
        three_line_table(document, headers, rows)
        table_records.append(dict(table_index=index, source_csv=source, columns=columns,
            source_row_indices=indices, percentage_columns=list(percents),
            signed_columns=list(signed), headers=headers, rendered_rows=rows))

    paragraph('CMERC 官方 Checkpoint 可靠性诊断报告', True)
    paragraph('本报告只分析作者公开的固定 checkpoint。完整输入直接复用已验证的逐话语预测；'
              '模态删除采用同一权重、官方执行副本及已记录兼容补丁的新推理。'
              '不重新训练、不选模、不拟合校准器；旧本地 Validation 六项实验独立保留。')
    paragraph('实验依据与可复现边界', True)
    paragraph('官方源码固定版本：815c7c060914f5b55db816c0ba3662b11d08552f。'
              'IEMOCAP_checkpoint.pkl 与 MELD_checkpoint.pkl 为纯 state_dict，严格加载；'
              '作者训练 seed、epoch 及选模协议缺少可核验元数据。源码存在测试集选模风险，'
              '不能据此断言公开权重的具体选模历史。')
    paragraph('运行保留官方 batch size=16、测试对话顺序、Padding/Mask 与默认 cuDNN RNN。'
              '兼容补丁仅删除已失效且未使用的 PyG topk 导入，并修正 HypergraphConv 第二次 '
              'propagate 的 size 轴顺序；不应描述为未经修改的官方运行环境。'
              '完整输入性能与协议详情见独立 official_checkpoint_verification 产物。')
    summary_name = 'official_cmerc_summary.csv'
    summary = sources[summary_name][1]
    table('表 1　完整输入分类与错误汇总（比例以 % 显示）', summary_name,
          ['dataset', 'valid_utterances', 'accuracy', 'weighted_f1', 'total_errors',
           'errors_ge_0p8', 'errors_ge_0p9', 'errors_ge_0p95'],
          ['数据集', '有效话语', 'Accuracy (%)', 'W-F1 (%)', '错误总数',
           '错误≥0.80', '错误≥0.90', '错误≥0.95'], percents=('accuracy', 'weighted_f1'))
    paragraph('实验一：完整输入高置信度错误', True)
    paragraph('置信度为预测类别最大概率 max P(class)；高置信度错误必须同时满足预测错误与'
              '置信度达到阈值。所有错误及各阈值的全部样本均保存在 CSV 与 Excel，'
              '不以典型案例替代全量统计。高置信度预测中的错误率分母为全部达到该阈值的预测数。')
    table('表 2　错误置信度与高置信度预测错误率', summary_name,
          ['dataset', 'mean_error_confidence', 'median_error_confidence',
           'high_confidence_predictions_ge_0p9', 'errors_ge_0p9',
           'high_confidence_error_rate_ge_0p9'],
          ['数据集', '错误平均置信度', '错误中位置信度', '预测≥0.90', '错误≥0.90',
           '高置信度预测错误率 (%)'], percents=('high_confidence_error_rate_ge_0p9',))
    paragraph('实验二：Baseline 与官方 CMERC 逐样本纠错比较', True)
    paragraph('Baseline comparison pending: matched checkpoint unavailable')
    paragraph('本报告未将旧 Validation Baseline 作为作者公开 CMERC 权重的严格配对对照。'
              '需要来源可追溯、数据特征与划分一致、模型及测试协议兼容，并具有选模记录的 '
              'M³Net Baseline checkpoint 后才能补齐实验二及其余 Baseline 对照。')
    if baseline_text:
        paragraph('Baseline 检索范围与证据边界', True)
        paragraph('已盘点可访问的本地和服务器 checkpoint，并检查官方 M³Net 仓库的 '
                  'README、源码树、标签及发布记录。未找到来源和评估协议均可验证的严格匹配权重，'
                  '不等于断言所有地方都不存在公开 Baseline。详细资源清单、链接、版本与检索限制保存在 '
                  'baseline_checkpoint_audit.md；该审计文件的 SHA-256 纳入本报告核验清单。')
        paragraph('本轮已实际取得 M³Net 官方分享的 ZIP，其中有三份 IEMOCAP 候选而无 MELD 权重。'
                  '预先固定 Roberta_IEMOCAP3.pth.tar，在 CPU 使用 weights_only=True 安全加载，'
                  '确认其为含 85 个张量的纯 OrderedDict；键和形状与官方 CMERC IEMOCAP 全部兼容，'
                  '可访问的四份特征文件哈希也相同。没有执行 Baseline 模型前向或比较候选测试分数。')
        paragraph('该候选 SHA-256：fb939de1f88855fda46a9d491e5cce4400c88c7187f3b792bb40f61c4ed03faa。'
                  '官方发布来源、结构和可用数据兼容已经证实；实际训练 seed、epoch、选模历史及'
                  '与作者 CMERC 的对应关系仍无法核验，因此尚不足以构成用户要求的严格配对。'
                  '权重保留在独立候选目录，不放入报告或 Git。')
    paragraph('实验三：残余错误与真实情绪类别分布', True)
    for dataset in DATASETS:
        table(f'表　{dataset} 按真实情绪的错误分布', f'{dataset}/class_error_summary.csv',
              ['true_emotion', 'support', 'total_errors', 'error_rate', 'errors_ge_0p8',
               'errors_ge_0p9', 'errors_ge_0p95'],
              ['真实情绪', '支持数', '错误总数', '类内错误率 (%)', '错误≥0.80',
               '错误≥0.90', '错误≥0.95'], percents=('error_rate',))
        paragraph(f'{dataset} 典型错误案例（按最大预测概率降序，展示前三条）', True)
        rows = sources[f'{dataset}/errors_all.csv'][1]
        for row in rows[:3]:
            paragraph(f'对话 {row["dialogue_id"]}，话语索引 {row["utterance_index"]}；'
                      f'真实情绪 {row["true_emotion"]}，预测 {row["predicted_emotion"]}；'
                      f'最大预测概率 {formatted(row["predicted_confidence"])}，'
                      f'真实类别概率 {formatted(row["true_label_probability"])}。')
            paragraph('原文：' + row['utterance_text'])
    paragraph('实验四：删除单一模态后的高置信度错误', True)
    paragraph('本实验复用官方 caldiff/cal_diff.py 的均值特征替换策略，保留其他模态及'
              '真实标签；具体特征轴、广播范围、batch 顺序与校验依据见 run_manifest 和'
              '模态删除协议审计。均值替换不等于置零，也不表示重新训练一个缺模态模型。')
    paragraph('特征形状为 [T, B, D]，沿时间轴 dim=0 求均值，再 repeat(T,1,1)。'
              '官方均值包含该 batch 的 Padding 时间步，不按有效 mask 重算；每个对话、特征维度'
              '分别保留自己的均值。删除文本同时替换四组文本特征，删除音频或视觉各只替换'
              '对应一组特征；其他输入、有效长度与标签保持原样。')
    for dataset in DATASETS:
        source = f'{dataset}/modality_ablation_summary.csv'
        fields = sources[source][0]
        condition = 'input_condition' if 'input_condition' in fields else 'removed_modality'
        table(f'表　{dataset} 模态删除汇总', source,
              [condition, 'total_errors', 'errors_ge_0p8', 'errors_ge_0p9',
               'errors_ge_0p95', 'mean_error_confidence'],
              ['输入条件', '错误总数', '错误≥0.80', '错误≥0.90', '错误≥0.95', '错误平均置信度'])
        table(f'表　{dataset} 模态删除后的分类性能', source,
              [condition, 'valid_samples', 'accuracy', 'weighted_f1'],
              ['输入条件', '有效话语', 'Accuracy (%)', 'W-F1 (%)'],
              percents=('accuracy', 'weighted_f1'))
        ablation_rows = sources[source][1]
        full_row = next(row for row in ablation_rows if row[condition] == 'full')
        for row in ablation_rows:
            if row[condition] == 'full':
                continue
            difference = int(row['total_errors']) - int(full_row['total_errors'])
            paragraph(f'{dataset} {row[condition]}：总错误 {row["total_errors"]}'
                      f'（相对 full {difference:+d}），W-F1={formatted(row["weighted_f1"], True)}%，'
                      f'置信度≥0.90 的错误 {row["errors_ge_0p9"]}。')
        paragraph('分类错误总数与 W-F1 必须和高置信度错误一起解读：若模型整体输出置信度较低，'
                  '删除模态仍可明显破坏分类性能，而高置信度错误数量保持为零；后者不表示删除无害。')
        transitions = [name for name in sources if name.startswith(dataset + '/') and
                       ('full_correct_to_masked' in name or 'transition_summary' in name)]
        for source in transitions:
            fields = sources[source][0]
            condition = 'input_condition' if 'input_condition' in fields else 'removed_modality'
            table(f'表　{dataset} 完整输入正确→删除后高置信度错误', source,
                  [condition, 'threshold', 'full_correct_samples', 'new_high_confidence_errors', 'ratio'],
                  ['输入条件', '阈值', '完整输入正确数', '新增高置信度错误', '占完整正确比例 (%)'],
                  percents=('ratio',))
            transition_rows = sources[source][1]
            seen, first_rows = {}, []
            for index, row in enumerate(transition_rows):
                key = row[condition]
                values = tuple(row[field] for field in
                               ('full_wrong_samples', 'still_wrong_count', 'still_wrong_ratio'))
                if key in seen and seen[key] != values:
                    raise ValueError('Threshold-independent still-wrong counts differ across threshold rows')
                if key not in seen:
                    seen[key] = values
                    first_rows.append(index)
            table(f'表　{dataset} 完整输入已错→删除后仍错', source,
                  [condition, 'full_wrong_samples', 'still_wrong_count', 'still_wrong_ratio'],
                  ['输入条件', '完整输入错误数', '删除后仍错数', '占完整错误比例 (%)'],
                  percents=('still_wrong_ratio',), selected=first_rows)
        paragraph('表注：完整输入正确→删除后高置信度错误的比例，以该模型完整输入正确话语数'
                  '为分母；完整输入已错→删除后仍错另列。最大错误类别概率与真实类别概率变化'
                  '是不同变量，不能将高置信度错误直接认定为模态冲突，也不能将置信度降低称为纠错。')
    paragraph('实验五：全部错误样本的置信度分布', True)
    for dataset in DATASETS:
        table(f'表　{dataset} 错误置信度固定分箱', f'{dataset}/error_confidence_distribution.csv',
              ['confidence_range', 'error_count', 'ratio_of_all_errors'],
              ['置信度区间', '错误数', '占全部错误比例 (%)'], percents=('ratio_of_all_errors',))
    paragraph('表注：各箱左闭右开，最后一箱 [0.95,1.0] 包含 1；箱数之和等于对应模型全部错误数。'
              '此处的比例分母是全部错误数，与高置信度预测中的错误率不同。')
    paragraph('实验六：完整测试集概率校准', True)
    paragraph('ECE/MCE 使用十个等宽箱 [0,0.1)、…、[0.9,1]。Signed Gap=平均置信度−准确率，'
              '正为过置信、负为置信度不足；空箱统计为 N/A，不绘制观测点。'
              'ECE=Σ(n箱/N)|Gap|，MCE=max非空箱|Gap|。多分类 Brier 对所有类别平方误差求和再'
              '按话语平均，不除以类别数或 2；NLL 使用真实类别概率的自然对数，不裁剪或重新归一化。')
    table('表　官方 CMERC 校准指标', summary_name,
          ['dataset', 'accuracy', 'mean_confidence', 'ece', 'mce', 'brier_score', 'nll'],
          ['数据集', 'Accuracy (%)', '平均置信度 (%)', 'ECE (%)', 'MCE (%)', 'Brier', 'NLL'],
          percents=('accuracy', 'mean_confidence', 'ece', 'mce'))
    for dataset in DATASETS:
        source = f'{dataset}/calibration_bins.csv'
        table(f'表　{dataset} 十等宽箱统计', source,
              ['lower_bound', 'upper_bound', 'count', 'mean_confidence', 'accuracy', 'signed_gap'],
              ['下界', '上界', '样本数', '平均置信度 (%)', '实际准确率 (%)', '差值 (pp)'],
              percents=('mean_confidence', 'accuracy', 'signed_gap'), signed=('signed_gap',))
        document.add_picture(str(directory / f'figures/reliability_{dataset}.png'), width=Inches(4.6))
        paragraph(f'{dataset} Reliability Diagram：虚线是理想校准，折线只连接非空箱的'
                  '真实平均置信度与准确率，下方显示十箱支持数。样本较少的箱具有较大描述性不稳定性。')
    paragraph('研究发现、局限与后续工作', True)
    for row in summary:
        paragraph(f'{row["dataset"]}：W-F1={formatted(row["weighted_f1"], True)}%，'
                  f'总错误 {row["total_errors"]}，其中置信度≥0.90 的错误 {row["errors_ge_0p9"]}；'
                  f'ECE={formatted(row["ece"], True)}%，MCE={formatted(row["mce"], True)}%，'
                  f'Brier={formatted(row["brier_score"])}，NLL={formatted(row["nll"])}。'
                  '达到论文分类分数仍然可以存在高置信度错误；分类性能与概率校准描述不同性质。')
        direction = '过度自信' if float(row['mean_confidence']) > float(row['accuracy']) else '置信度不足'
        paragraph(f'{row["dataset"]} 平均置信度 {formatted(row["mean_confidence"], True)}%，'
                  f'实际准确率 {formatted(row["accuracy"], True)}%，整体均值呈{direction}。'
                  '逐箱偏差仍需参照可靠性图和分箱支持数，均值差不能代替 ECE。')
        if int(row['high_confidence_predictions_ge_0p9']) == 0:
            paragraph(f'{row["dataset"]} 没有置信度≥0.90 的预测，其高置信度预测错误率'
                      '分母为零，必须为 N/A。高置信度错误数为零不能解释为模型完美校准；'
                      '模型仍存在大量较低置信度错误，且可能明显低估预测正确率。')
    paragraph('当前分析是固定权重的测试集描述性诊断，不能证明跨随机种子稳定性、统计显著性或'
              '错误的因果机制。缺少严格匹配 Baseline 时，不能声称官方 CMERC 对 Baseline 的'
              '纠错、校准或删除模态鲁棒性改善。后续应先补齐可追溯配对权重，再进行同协议对照。')
    # Make the font explicit on every run, including picture paragraphs.
    for paragraph_ in document.paragraphs:
        for run in paragraph_.runs:
            _set_run(run)
    document.save(directory / DOCX_NAME)
    return table_records


def make_excel(path, sources):
    records = []
    aliases = {'all_predictions': '全预测', 'errors_all': '全部错误',
        'high_confidence_errors': '高置信错误', 'class_error_summary': '类别错误',
        'error_confidence_distribution': '错误分布', 'calibration_bins': '校准分箱',
        'modality_ablation_predictions': '删除预测', 'modality_ablation_summary': '删除汇总',
        'full_correct_to_masked_error': '删除新增错误', 'modality_transition_summary': '预测转移',
        'official_cmerc_summary': '官方总体统计'}
    with xlsxwriter.Workbook(str(path), {'strings_to_formulas': False,
                                       'strings_to_urls': False}) as workbook:
        head = workbook.add_format({'font_name': FONT, 'bold': True,
                                    'text_wrap': True, 'bg_color': '#E8EDF1'})
        body = workbook.add_format({'font_name': FONT, 'num_format': '0.000000000000000',
                                    'valign': 'top'})
        info = workbook.add_worksheet('说明')
        notes = [
            '作者公开固定 checkpoint 诊断；未训练、未拟合校准、未混用旧 Validation 权重。',
            '所有比例与概率保留源 CSV 的 0–1 单位；Accuracy/W-F1 也是 0–1，Word 表中另转为 %。',
            'ECE/MCE：10 等宽箱；空箱的均值、准确率、差值为空或 N/A。Signed Gap=置信度−准确率。',
            'Brier：各类别平方误差之和的样本均值；NLL 使用自然对数，不裁剪概率。',
            'high_confidence_errors 按阈值展开，同一句可出现在多个阈值下，使用样本 ID 加 threshold 唯一标识。',
            '高置信度预测错误率分母是所有达到阈值的预测；错误分布占比分母是全部错误。',
            '删除后新增高置信度错误比例，分母为完整输入预测正确的话语数。',
            'Baseline comparison pending: matched checkpoint unavailable',
            '所有 CSV 完整保留；本工作簿按数据集、实验归档全量逐话语结果，未截取 Top 10。',
        ]
        for index, note in enumerate(notes):
            info.write_string(index, 0, note, body)
        info.set_column(0, 0, 135)
        used = {'说明'}
        for source, (fields, rows) in sources.items():
            parts = Path(source).parts
            dataset = parts[0] + '_' if len(parts) > 1 else ''
            stem = Path(source).stem
            name = (dataset + aliases.get(stem, stem))[:31]
            if name in used:
                raise ValueError(f'Duplicate worksheet name: {name}')
            used.add(name)
            sheet = workbook.add_worksheet(name)
            sheet.write_row(0, 0, fields, head)
            for i, row in enumerate(rows, 1):
                for j, field in enumerate(fields):
                    value = excel_value(field, row[field])
                    result = (sheet.write_number(i, j, value, body) if isinstance(value, (int, float))
                              else sheet.write_string(i, j, value, body))
                    if result != 0:
                        raise ValueError(f'Excel cell truncated or invalid: {source}:{i}:{field}')
            sheet.freeze_panes(1, 0)
            sheet.autofilter(0, 0, len(rows), len(fields) - 1)
            for j, field in enumerate(fields):
                sheet.set_column(j, j, 55 if field == 'utterance_text' else min(32, max(15, len(field))))
            records.append(dict(name=name, worksheet_index=len(records) + 2,
                                source_csv=source, columns=fields, row_count=len(rows)))
    return records


def verify_excel(path, sources, records):
    checks, cells_checked = [], 0
    with zipfile.ZipFile(path) as archive:
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            root = ET.fromstring(archive.read('xl/sharedStrings.xml'))
            strings = [''.join(node.itertext()) for node in root.findall('s:si', NS)]
        for record in records:
            root = ET.fromstring(archive.read(f'xl/worksheets/sheet{record["worksheet_index"]}.xml'))
            xml_rows = root.findall('s:sheetData/s:row', NS)
            fields, rows = sources[record['source_csv']]
            expected_rows = [fields] + [[excel_value(f, r[f]) for f in fields] for r in rows]
            if len(xml_rows) != len(expected_rows):
                raise AssertionError(f'Excel row count mismatch: {record["name"]}')
            for xml_row, expected in zip(xml_rows, expected_rows):
                cells = xml_row.findall('s:c', NS)
                if len(cells) != len(expected):
                    raise AssertionError(f'Excel column mismatch: {record["name"]}')
                for cell, value in zip(cells, expected):
                    node = cell.find('s:v', NS)
                    actual = '' if node is None else node.text or ''
                    if cell.get('t') == 's':
                        actual = strings[int(actual)]
                    elif cell.get('t') == 'inlineStr':
                        actual = ''.join(cell.find('s:is', NS).itertext())
                    if isinstance(value, (int, float)):
                        if not math.isclose(float(actual), value, rel_tol=2e-15, abs_tol=1e-15):
                            raise AssertionError(f'Excel numeric mismatch: {record["name"]}, {value}, {actual}')
                    elif actual != value:
                        raise AssertionError(f'Excel text mismatch: {record["name"]}')
                    cells_checked += 1
            checks.append(dict(**record, all_cells_match_csv=True))
    return dict(passed=True, cells_checked=cells_checked, worksheets=checks,
                numeric_tolerance='relative 2e-15 or absolute 1e-15; Excel binary floating point')


def verify_word(path, tables):
    saved = Document(path)
    for record in tables:
        cells = [[cell.text for cell in row.cells] for row in saved.tables[record['table_index']].rows]
        if cells != [record['headers']] + record['rendered_rows']:
            raise AssertionError('Word table values differ from CSV display transformation')
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
        if any(name.startswith(('word/header', 'word/footer')) for name in archive.namelist()):
            raise AssertionError('Unexpected header/footer files')
        if root.findall('.//w:headerReference', NS) or root.findall('.//w:footerReference', NS):
            raise AssertionError('Unexpected header/footer references')
        runs = root.findall('.//w:r', NS)
        for run in runs:
            fonts = run.find('w:rPr/w:rFonts', NS)
            color = run.find('w:rPr/w:color', NS)
            size = run.find('w:rPr/w:sz', NS)
            if (fonts is None or any(fonts.get(qn(f'w:{slot}')) != FONT
                                     for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs')) or
                color is None or color.get(qn('w:val')) != '000000' or
                size is None or size.get(qn('w:val')) != '24'):
                raise AssertionError('Word run font/color/12pt mismatch')
        for table in root.findall('.//w:tbl', NS):
            borders = table.find('w:tblPr/w:tblBorders', NS)
            if borders is None:
                raise AssertionError('Word table lacks explicit three-line borders')
            for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
                border = borders.find(f'w:{side}', NS)
                expected = 'single' if side in ('top', 'bottom') else 'nil'
                if border is None or border.get(qn('w:val')) != expected:
                    raise AssertionError('Word table border rule mismatch')
            for cell in table.findall('w:tr', NS)[0].findall('w:tc', NS):
                edge = cell.find('w:tcPr/w:tcBorders/w:bottom', NS)
                if edge is None or edge.get(qn('w:val')) != 'single':
                    raise AssertionError('Missing third line beneath Word header')
        pictures = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/picture}pic')
        if len(pictures) != 2:
            raise AssertionError('Expected two reliability images')
    return dict(passed=True, tables_checked=len(tables), all_table_cells_match_csv=True,
                font=FONT, font_size_points=12, color='000000', runs_checked=len(runs),
                explicit_three_line_tables=True, reliability_images=2, headers_footers=False,
                font_embedding=False, office_layout_rendered=False,
                layout_note='OOXML formatting and numbers verified; Office/LibreOffice pagination not rendered.',
                table_provenance=tables)


def render(output_dir):
    output = Path(output_dir).resolve()
    required = ['official_cmerc_summary.csv', 'baseline_checkpoint_audit.md'] + [
        f'{dataset}/{name}.csv' for dataset in DATASETS for name in (
            'all_predictions', 'errors_all', 'high_confidence_errors', 'class_error_summary',
            'modality_ablation_predictions', 'modality_ablation_summary',
            'error_confidence_distribution', 'calibration_bins')]
    for name in required:
        if not (output / name).is_file():
            raise FileNotFoundError(output / name)
    # Deliberately collect only this independent run, never the parent Validation outputs.
    paths = sorted(output.glob('*.csv')) + [path for dataset in DATASETS
                                          for path in sorted((output / dataset).glob('*.csv'))]
    sources = {str(path.relative_to(output)): read_csv(path) for path in paths}
    source_hashes = {str(path.relative_to(output)): digest(path) for path in paths}
    baseline_path = output / 'baseline_checkpoint_audit.md'
    source_hashes[baseline_path.name] = digest(baseline_path)
    baseline_text = baseline_path.read_text(encoding='utf-8')
    if 'Baseline comparison pending: matched checkpoint unavailable' not in baseline_text:
        raise ValueError('Renderer currently requires the audited matched-Baseline-unavailable result')
    groups = {dataset: sorted(sources[f'{dataset}/calibration_bins.csv'][1],
                              key=lambda row: int(row['bin_index'])) for dataset in DATASETS}
    for dataset, bins in groups.items():
        if [int(row['bin_index']) for row in bins] != list(range(10)):
            raise ValueError('Missing calibration bins')
        expected = 1623 if dataset == 'IEMOCAP' else 2610
        if sum(int(row['count']) for row in bins) != expected:
            raise ValueError('Calibration support differs from verified official full predictions')
    names = [DOCX_NAME, XLSX_NAME, 'report_verification.json', 'figure_contract.md'] + [
        f'figures/reliability_{dataset}.{extension}' for dataset in DATASETS
        for extension in ('png', 'svg', 'pdf')]
    for name in names:
        if (output / name).exists():
            raise FileExistsError(f'Refusing to overwrite report: {output / name}')
    with tempfile.TemporaryDirectory(prefix='.official_render_', dir=output) as folder:
        temporary = Path(folder)
        (temporary / 'figures').mkdir()
        (temporary / 'figure_contract.md').write_text(figure_contract(groups), encoding='utf-8')
        figures = make_figures(temporary, groups)
        tables = make_word(temporary, sources, baseline_text)
        sheets = make_excel(temporary / XLSX_NAME, sources)
        verification = dict(status='passed', training=False, gpu_inference=False,
            previous_validation_results_used=False, sources=source_hashes,
            word=verify_word(temporary / DOCX_NAME, tables),
            excel=verify_excel(temporary / XLSX_NAME, sources, sheets), figures=figures,
            outputs={name: dict(sha256=digest(temporary / name),
                               bytes=(temporary / name).stat().st_size)
                     for name in names if name != 'report_verification.json'})
        for name, expected in source_hashes.items():
            if digest(output / name) != expected:
                raise RuntimeError(f'Source changed while rendering: {name}')
        (temporary / 'report_verification.json').write_text(
            json.dumps(verification, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        installed = []
        try:
            for name in names:
                (output / name).parent.mkdir(parents=True, exist_ok=True)
                os.link(temporary / name, output / name)
                installed.append(output / name)
        except Exception:
            for path in installed:
                path.unlink()
            raise
    return verification


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    result = render(args.output_dir)
    print(json.dumps({'status': result['status'], 'word_tables': result['word']['tables_checked'],
                      'excel_sheets': len(result['excel']['worksheets']),
                      'excel_cells_checked': result['excel']['cells_checked']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
