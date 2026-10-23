"""Synthetic calibration report QA; no model inference or original files edited."""
import csv
import hashlib
import importlib
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
import pytest


def renderer():
    return importlib.import_module('scripts.render_calibration_experiment')


def sources(tmp_path):
    out = tmp_path / 'exp6'
    out.mkdir()
    mother = tmp_path / 'previous.docx'
    doc = Document()
    doc.add_paragraph('原实验一至五正文')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '旧表格'
    doc.save(mother)
    overall, bins = [], []
    for dataset in ('IEMOCAP', 'MELD'):
        for model in ('baseline', 'cmerc'):
            overall.append(dict(dataset=dataset, model=model, input_condition='full', valid_utterances=2,
                                accuracy=.5, mean_confidence=.7, ece=.2, mce=.2,
                                brier_score=.6, nll=1.2, n_bins=10, zero_true_probability_count=0))
            for i in range(10):
                bins.append(dict(dataset=dataset, model=model, bin_index=i,
                                 lower_bound=i/10, upper_bound=(i+1)/10, count=2 if i==7 else 0,
                                 mean_confidence=.7 if i==7 else '', accuracy=.5 if i==7 else '',
                                 signed_gap=.2 if i==7 else '', absolute_gap=.2 if i==7 else '',
                                 weight=1 if i==7 else 0, ece_contribution=.2 if i==7 else 0))
    for filename, data in (('exp6_overall_summary.csv', overall), ('exp6_confidence_bins.csv', bins)):
        with (out / filename).open('w', encoding='utf-8', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(data[0]))
            writer.writeheader()
            writer.writerows(data)
    (out / 'SUMMARY.md').write_text('## 主要发现\n此处是真实统计结论占位测试。\n\n| repeated | table |\n|---|---|\n| a | b |\n', encoding='utf-8')
    (out / 'run_manifest.json').write_text('{}\n', encoding='utf-8')
    return out, mother


def test_figures_word_excel_and_provenance(tmp_path):
    out, mother = sources(tmp_path)
    mod = renderer()
    before = mother.read_bytes()
    prefix = mod.body_prefix(Document(mother))
    section_xml = Document(mother).element.body.sectPr.xml
    input_bytes = {name:(out/name).read_bytes() for name in mod.INPUT_FILES}
    manifest = mod.render_calibration_reports(out, mother)
    assert mother.read_bytes() == before
    assert manifest['mother_document']['sha256'] == hashlib.sha256(before).hexdigest()
    assert manifest['mother_document']['body_prefix_preserved']
    new = Document(out / mod.DOCX_NAME)
    assert new.element.body.sectPr.xml == section_xml
    assert all((out/name).read_bytes() == data for name,data in input_bytes.items())
    assert mod.body_prefix(new)[:len(prefix)] == prefix
    assert len(manifest['word_tables']) == 5
    assert len(new.inline_shapes) == 2
    for table in manifest['word_tables']:
        actual = [[c.text for c in row.cells] for row in new.tables[table['document_table_index']].rows]
        assert mod.string_hash(actual) == table['cells_sha256']
        assert len(actual)-1 == (4 if table['kind']=='overall' else 10)
        xml = new.tables[table['document_table_index']]._tbl.xml
        assert 'w:val="single"' in xml
        assert 'w:left w:val="nil"' in xml
        if table['kind'] == 'bins':
            assert actual[0] == ['Confidence interval','Count','Mean confidence (%)','Accuracy (%)','Signed gap (pp)']
            assert actual[8][2:] == ['70.0000','50.0000','+20.000000']
            assert actual[1][2:] == ['N/A','N/A','N/A']
            assert 'multiplied by 100' in table['display_rule']
    overall_table = new.tables[manifest['word_tables'][0]['document_table_index']]
    assert overall_table.cell(0, 2).text == 'ECE (%)'
    assert overall_table.cell(1, 2).text == '20.0000'
    assert (out / 'figure_contract.md').exists()
    for figure in manifest['figures']:
        for fmt in ('png', 'pdf', 'svg'):
            assert Path(figure['files'][fmt]['path']).is_file()
        svg = Path(figure['files']['svg']['path']).read_text(encoding='utf-8')
        assert '<text' in svg
        for curve in figure['curves']:
            assert curve['bin_indices'] == [7]
            assert curve['points'] == [[.7, .5]]
            assert len(curve['counts']) == 10
        assert figure['points']['baseline'] == [dict(bin_index=7,count=2,source_row_index=(7 if figure['dataset']=='IEMOCAP' else 27),x=.7,y=.5)]
    with ZipFile(out / mod.XLSX_NAME) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        assert len(workbook.findall('s:sheets/s:sheet', ns)) == 5
    assert not any(table.cell(0,0).text == 'repeated' for table in new.tables)


def test_no_overwrite_and_missing_mother(tmp_path):
    out, mother = sources(tmp_path)
    with pytest.raises(FileNotFoundError):
        renderer().render_calibration_reports(out, out/'missing.docx')
    sentinel = out / renderer().DOCX_NAME
    sentinel.write_bytes(b'keep')
    with pytest.raises(FileExistsError):
        renderer().render_calibration_reports(out, mother)
    assert sentinel.read_bytes() == b'keep'
    assert not list(out.glob('*.png'))


def test_empty_bin_must_not_have_fabricated_point(tmp_path):
    out, mother = sources(tmp_path)
    source = out / 'exp6_confidence_bins.csv'
    source.write_text(source.read_text().replace('0.0,0.1,0,,', '0.0,0.1,0,0,'), encoding='utf-8')
    with pytest.raises(ValueError, match='empty bin'):
        renderer().render_calibration_reports(out, mother)
