"""Render experiment 2–5 CSVs without modifying experiment 1 or model artifacts.

No inference or checkpoint selection occurs here. All table values are read from
the supplied CSVs. The output manifest makes CSV, Excel and Word provenance
independently auditable without requiring an Excel reader dependency.
"""
import argparse
from collections import defaultdict
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from lxml import etree
import xlsxwriter


DOCX_NAME = 'CMERC_高置信度错误诊断报告_完整版.docx'
XLSX_NAME = 'CMERC_高置信度错误诊断结果_完整版.xlsx'
MOTHER_NAME = 'CMERC_实验一_高置信度错误样本统计表.docx'
SUMMARY_TABLES = {
    2: ('exp2_summary.csv',),
    3: ('exp3_class_summary.csv',),
    4: ('exp4_modality_ablation_summary.csv', 'exp4_full_correct_to_masked_error.csv'),
    5: ('exp5_error_confidence_distribution.csv', 'exp5_overall_summary.csv'),
}
TITLES = {
    2: '实验二：CMERC 对 Baseline 高置信度错误的改善情况',
    3: '实验三：CMERC 校准后仍然存在的高置信度错误',
    4: '实验四：删除单一模态后的高置信度错误',
    5: '实验五：Baseline 与 CMERC 错误样本的置信度分布',
}
NOTES = {
    2: '阈值为 0.80、0.90、0.95。corrected 才表示分类得到纠正；'
       'still_wrong_confidence_lower 表示仍然错误但最大置信度降低。'
       '改善比例的分母为 Baseline 在对应阈值下的高置信度错误数；'
       '新错误比例的分母见表中 eligible 字段，不能与纠错比例混用。'
       '两个模型预测错误类别不同时，最大置信度降低不代表同一错误类别概率降低。',
    3: '类内错误率的分母为该真实情绪类别的全部有效话语；'
       '高置信度预测中的错误率分母为该真实类别内达到阈值的全部预测。'
       '各阈值集合相互包含，不能直接相加。逐话语全量样本保存在 CSV 和 Excel。',
    4: 'full 为完整三模态；w/o_text、w/o_audio、w/o_visual 为对应模态均值替换。'
       '完整输入正确后新增高置信度错误比例的分母为完整输入正确的话语数。'
       '完整输入本来就错误且删除后仍错误的样本另行统计。'
       '错误预测的最大概率与真实类别概率上升（以及 CEL）是不同量，不能混用。',
    5: '置信度区间左闭右开，最后一个区间 [0.95, 1.0] 包含 1。'
       '区间内错误占比的分母为该模型全部错误数；'
       '高置信度预测中的错误率分母为达到该阈值的全部预测数。'
       '分母为零时显示 N/A，不能解释成错误率为零。',
}


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def string_hash(rows):
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False, separators=(',', ':')).encode('utf-8')).hexdigest()


def display(value):
    """Keep original text intact; only explicit missing-value tokens become N/A."""
    if value is None or (isinstance(value, float) and value != value):
        return 'N/A'
    text = str(value)
    if text in ('', 'None', 'null', 'NaN', 'nan', 'N/A'):
        return 'N/A'
    return text


def validate_text(value, location):
    # Word XML cannot represent these characters. Fail explicitly rather than
    # silently dropping them or altering the original transcript / source CSV.
    for char in value:
        code = ord(char)
        if ((code < 32 and char not in '\t\n\r') or 0xD800 <= code <= 0xDFFF
                or code in (0xFFFE, 0xFFFF)):
            raise ValueError(f'Unsupported XML control character U+{code:04X}: {location}')


def body_prefix(document):
    """Canonical original body elements, excluding its terminal section metadata."""
    return [etree.tostring(child, method='c14n') for child in document.element.body
            if etree.QName(child).localname != 'sectPr']


def _read_sources(output_dir):
    csv.field_size_limit(16 * 1024 * 1024)
    paths = sorted(output_dir.glob('exp*.csv'))
    full = output_dir / 'full_predictions_reused.csv'
    if full.exists():
        paths.append(full)
    required = {name for names in SUMMARY_TABLES.values() for name in names}
    missing = required - {path.name for path in paths}
    if missing:
        raise FileNotFoundError(f'Missing summary CSVs: {sorted(missing)}')
    sources = []
    for path in paths:
        source_hash = sha256(path)
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            headers = reader.fieldnames
            if not headers or len(headers) != len(set(headers)):
                raise ValueError(f'Missing or duplicate CSV headers: {path}')
            for header in headers:
                validate_text(header, str(path))
            rows = []
            for index, row in enumerate(reader):
                if None in row or any(value is None for value in row.values()):
                    raise ValueError(f'Malformed CSV row {index}: {path}')
                for key, value in row.items():
                    validate_text(value, f'{path} row {index}, {key}')
                rows.append(row)
        groups = defaultdict(list)
        dataset_key = next((key for key in headers if key.lower() == 'dataset'), None)
        if dataset_key:
            # An absent dataset in an error-only table means zero qualifying
            # rows, not missing coverage. Keep explicit header-only sheets.
            groups['IEMOCAP'] = []
            groups['MELD'] = []
        for index, row in enumerate(rows):
            groups[row.get(dataset_key, 'ALL') or 'ALL'].append(index)
        if not rows and not dataset_key:
            groups['EMPTY'] = []
        sources.append(dict(path=path.resolve(), sha256=source_hash,
                            headers=headers, rows=rows, groups=dict(groups)))
    return sources


def _sheet_name(stem, dataset, used):
    short = stem.replace('modality_ablation', 'mod_ablat').replace('high_confidence', 'hc')
    short = short.replace('confidence', 'conf').replace('predictions', 'pred')
    base = re.sub(r'[\[\]:*?/\\]', '_', f'{dataset}_{short}').strip("'")
    candidate = base[:31]
    suffix = 1
    while candidate.lower() in used:
        suffix += 1
        candidate = base[:31-len(str(suffix))-1] + '_' + str(suffix)
    used.add(candidate.lower())
    return candidate


def _xlsx(path, sources):
    sheets = []
    workbook = xlsxwriter.Workbook(str(path), {'strings_to_formulas': False,
                                               'strings_to_urls': False,
                                               'strings_to_numbers': False})
    try:
        header_format = workbook.add_format({'bold': True, 'bg_color': '#DDEBF7', 'text_wrap': True, 'valign': 'top'})
        cell_format = workbook.add_format({'valign': 'top', 'text_wrap': True})
        index_sheet = workbook.add_worksheet('来源索引')
        used = {'来源索引'.lower()}
        index_headers = ['CSV', 'Source Path', 'SHA-256', 'Dataset', 'CSV Total Rows',
                         'Sheet', 'Sheet Rows', 'Source Row Indices SHA-256', 'Rendered Rows SHA-256']
        index_sheet.write_row(0, 0, index_headers, header_format)
        for source in sources:
            for dataset, indices in source['groups'].items():
                name = _sheet_name(source['path'].stem, dataset, used)
                worksheet = workbook.add_worksheet(name)
                worksheet.freeze_panes(1, 0)
                headers = source['headers']
                worksheet.write_row(0, 0, headers, header_format)
                rendered = []
                for target_index, source_index in enumerate(indices, 1):
                    row = [display(source['rows'][source_index][key]) for key in headers]
                    rendered.append(row)
                    for column, value in enumerate(row):
                        if len(value) > 32767:
                            raise ValueError(f'Excel cell exceeds 32767 characters: {source["path"]}, row {source_index}')
                        result = worksheet.write_string(target_index, column, value, cell_format)
                        if result != 0:
                            raise ValueError(f'Excel rejected cell ({target_index}, {column}), code {result}')
                worksheet.autofilter(0, 0, len(indices), len(headers)-1)
                for column, header in enumerate(headers):
                    width = 65 if ('text' in header.lower() or 'probabilities' in header.lower()) else min(32, max(14, len(header)+2))
                    worksheet.set_column(column, column, width)
                entry = dict(name=name, dataset=dataset, source_csv=str(source['path']),
                             columns=headers, row_count=len(indices), source_row_indices=indices,
                             rendered_rows_sha256=string_hash(rendered), worksheet_index=len(sheets)+2)
                sheets.append(entry)
                index_values = [source['path'].name, str(source['path']), source['sha256'],
                                dataset, str(len(source['rows'])), name, str(len(indices)),
                                string_hash(indices), entry['rendered_rows_sha256']]
                for column, value in enumerate(index_values):
                    index_sheet.write_string(len(sheets), column, value, cell_format)
        index_sheet.freeze_panes(1, 0)
        index_sheet.autofilter(0, 0, len(sheets), len(index_headers)-1)
        index_sheet.set_column(0, 0, 48)
        index_sheet.set_column(1, 2, 70)
        index_sheet.set_column(3, 6, 22)
        index_sheet.set_column(7, 8, 70)
    finally:
        workbook.close()
    return sheets


def _word_table(document, headers, rows, font_size=7):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    repeat.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'true')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    return table


def _append_markdown(document, text):
    """Transcribe source Markdown paragraphs and tables; do not edit its source."""
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.lstrip().startswith('|') and index+1 < len(lines) and re.fullmatch(r'[\s|:\-]+', lines[index+1]):
            parse = lambda item: [part.strip() for part in item.strip().strip('|').split('|')]
            headers = parse(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].lstrip().startswith('|'):
                values = parse(lines[index])
                if len(values) != len(headers):
                    break
                rows.append(values)
                index += 1
            _word_table(document, headers, rows, 8)
            continue
        if line:
            heading = re.match(r'^(#{1,6})\s+(.*)', line)
            if heading:
                document.add_heading(heading.group(2), min(len(heading.group(1)), 4))
            else:
                document.add_paragraph(line)
        else:
            document.add_paragraph('')
        index += 1


def _landscape(document):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.7)
    section.page_height = Inches(8.3)
    section.left_margin = section.right_margin = Inches(0.45)
    section.top_margin = section.bottom_margin = Inches(0.5)


def _column_groups(headers, maximum=10):
    """Repeat row identifiers, keep threshold triples and count/ratio pairs.

    A source summary can have many metrics. Split only its columns, never its
    rows, so each small table remains auditable against the complete source.
    """
    identifiers = [key for key in headers if key in (
        'dataset', 'model', 'threshold', 'input_condition', 'removed_modality',
        'true_label', 'true_emotion', 'confidence_bin')]
    if len(headers) <= maximum:
        return [('完整统计', headers)]
    capacity = maximum - len(identifiers)
    if capacity < 1:
        raise ValueError('Too many identifier columns for readable Word summary tables')
    blocks = defaultdict(list)
    for key in headers:
        if key in identifiers:
            continue
        threshold = re.search(r'_ge_(0p\d+)$', key)
        label = f'阈值 {threshold.group(1).replace("p", ".")}' if threshold else '基础统计'
        blocks[label].append(key)
    result = []
    for label, metrics in blocks.items():
        units, index = [], 0
        while index < len(metrics):
            unit = [metrics[index]]
            if (metrics[index].endswith('_count') and index+1 < len(metrics)
                    and metrics[index+1] == metrics[index][:-6]+'_ratio'):
                unit.append(metrics[index+1])
            units.append(unit)
            index += len(unit)
        chunk, part = [], 1
        for unit in units:
            if len(unit) > capacity:
                raise ValueError('A linked metric group cannot fit in a Word summary table')
            if chunk and len(chunk)+len(unit) > capacity:
                result.append((f'{label}（{part}）', identifiers+chunk))
                chunk, part = [], part+1
            chunk.extend(unit)
        if chunk:
            result.append((f'{label}（{part}）', identifiers+chunk))
    if set(key for _, group in result for key in group) != set(headers):
        raise AssertionError('Summary column splitting lost source fields')
    return result


def _narrative_markdown(text):
    """Exclude repeated Markdown tables from the appended findings section."""
    lines, output, index = text.splitlines(), [], 0
    while index < len(lines):
        if (lines[index].lstrip().startswith('|') and index+1 < len(lines)
                and re.fullmatch(r'[\s|:\-]+', lines[index+1])):
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith('|'):
                index += 1
            continue
        output.append(lines[index])
        index += 1
    return '\n'.join(output)


def _word(path, sources, exp1_root, mother, allow_missing_mother, output_dir):
    source_summary = exp1_root / 'SUMMARY.md'
    original_markdown = source_summary.read_text(encoding='utf-8')
    validate_text(original_markdown, str(source_summary))
    original_sha = sha256(mother) if mother else None
    original_body = None
    if mother:
        document = Document(mother)
        original_body = body_prefix(document)
    else:
        if not allow_missing_mother:
            raise FileNotFoundError('未找到实验一 Word 母版；需要显式 --allow-missing-mother 才能转录原 SUMMARY.md')
        document = Document()
        document.add_heading('CMERC 高置信度错误诊断报告（完整版）', 0)
        document.add_paragraph('未找到指定原Word母版，实验一内容从原SUMMARY.md转录，不能宣称继承母版版式。原始实验一文件保持不变。')
        document.add_paragraph(f'实验一摘要来源：{source_summary.resolve()}；SHA-256：{sha256(source_summary)}')
        _append_markdown(document, original_markdown)
    tables = []
    by_name = {source['path'].name: source for source in sources}
    for experiment, filenames in SUMMARY_TABLES.items():
        _landscape(document)
        document.add_heading(TITLES[experiment], 1)
        document.add_paragraph(NOTES[experiment])
        for filename in filenames:
            source = by_name[filename]
            for dataset, indices in source['groups'].items():
                document.add_heading(f'{dataset} — {filename}', 2)
                for group_label, headers in _column_groups(source['headers']):
                    document.add_paragraph(group_label)
                    rows = [[display(source['rows'][i][key]) for key in headers] for i in indices]
                    table_index = len(document.tables)
                    _word_table(document, headers, rows, 8)
                    if not rows:
                        document.add_paragraph('该表无符合条件的数据行；空表不能解释为模型完美。')
                    document.add_paragraph('表注：同一来源的分组表重复标识列，保留全部源字段和全部数据行。保持源 CSV 数值字符串，比例字段是 0–1 比例而非自动换算的百分数；N/A 表示缺失或分母为零。实际分母及阈值见表中字段和本节说明。')
                    tables.append(dict(experiment=experiment, dataset=dataset, source_csv=str(source['path']),
                                       column_group=group_label, columns=headers,
                                       source_row_indices=indices, row_count=len(rows),
                                       document_table_index=table_index, rendered_rows_sha256=string_hash(rows)))
    _landscape(document)
    document.add_heading('后续实验总结、研究发现、局限性与下一步计划', 1)
    summary = output_dir / 'SUMMARY.md'
    if summary.exists():
        new_summary = summary.read_text(encoding='utf-8')
        if not new_summary.startswith(original_markdown):
            raise ValueError('Follow-up SUMMARY.md must preserve experiment 1 SUMMARY.md as its complete prefix')
        _append_markdown(document, _narrative_markdown(new_summary[len(original_markdown):]))
    document.add_paragraph('所有结论仅适用于清单中已固定 checkpoint、相同测试样本及本次推理协议。未训练、未用测试集选择 checkpoint 或调参。历史 checkpoint 是否曾用测试集选模，以运行清单和原实验记录中的风险说明为准。')
    document.add_paragraph('置信度下降不表示错误已经纠正；高置信度错误不能直接解释成模态冲突。均值替换后的错误最大概率与真实类别概率变化须分开分析。后续可在独立验证集设计校准和机制验证，不能用本测试结果反复调参。')
    if original_body is not None and body_prefix(document)[:len(original_body)] != original_body:
        raise AssertionError('Original Word body changed during append')
    document.save(path)
    saved = Document(path)
    if original_body is not None and body_prefix(saved)[:len(original_body)] != original_body:
        raise AssertionError('Original Word body changed after serialization')
    if mother and sha256(mother) != original_sha:
        raise AssertionError('Mother document changed')
    for entry in tables:
        actual = [[cell.text for cell in row.cells] for row in saved.tables[entry['document_table_index']].rows]
        if actual[0] != entry['columns'] or string_hash(actual[1:]) != entry['rendered_rows_sha256']:
            raise AssertionError(f'Word table differs from source: {entry["source_csv"]}')
    return tables, dict(found=mother is not None, path=str(mother) if mother else None,
                        sha256=original_sha, body_prefix_preserved=True if mother else None,
                        original_body_element_count=len(original_body) if mother else None,
                        fallback_source=str(source_summary.resolve()), fallback_source_sha256=sha256(source_summary))


def render_reports(exp1_root, output_dir, mother_docx=None, allow_missing_mother=False):
    exp1_root, output_dir = Path(exp1_root).resolve(), Path(output_dir).resolve()
    if not output_dir.is_dir():
        raise FileNotFoundError(output_dir)
    names = (DOCX_NAME, XLSX_NAME, 'report_manifest.json')
    for name in names:
        if (output_dir / name).exists():
            raise FileExistsError(f'Refusing to overwrite existing report: {output_dir / name}')
    if mother_docx:
        mother = Path(mother_docx).resolve()
        if not mother.exists():
            if not allow_missing_mother:
                raise FileNotFoundError(f'实验一 Word 母版不存在：{mother}')
            mother = None
    else:
        candidates = sorted(exp1_root.rglob(MOTHER_NAME))
        if len(candidates) > 1:
            raise ValueError('多个实验一 Word 母版，请通过 --mother-docx 指定')
        mother = candidates[0] if candidates else None
    if mother is None and not allow_missing_mother:
        raise FileNotFoundError('未找到实验一 Word 母版，请指定 --mother-docx 或显式 --allow-missing-mother')
    sources = _read_sources(output_dir)
    with tempfile.TemporaryDirectory(prefix='.render_reports_', dir=output_dir) as directory:
        temporary = Path(directory)
        sheets = _xlsx(temporary / XLSX_NAME, sources)
        tables, provenance = _word(temporary / DOCX_NAME, sources, exp1_root, mother,
                                   allow_missing_mother, output_dir)
        manifest = dict(schema_version=1, mother_document=provenance, sheets=sheets, word_tables=tables,
                        sources=[dict(path=str(s['path']), sha256=s['sha256'], row_count=len(s['rows']),
                                      columns=s['headers']) for s in sources],
                        outputs={name: dict(path=str(output_dir/name), sha256=sha256(temporary/name))
                                 for name in (DOCX_NAME, XLSX_NAME)},
                        notes=['Source data rows are indexed from zero, excluding CSV header.',
                               'All source rows are included in Excel dataset sheets.',
                               'All Excel cells are strings; formulas and URLs are not interpreted.',
                               'Missing-value tokens are rendered N/A; original CSVs are untouched.'])
        for source in sources:
            if sha256(source['path']) != source['sha256']:
                raise RuntimeError(f'Source CSV changed during rendering: {source["path"]}')
        (temporary / 'report_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        # Hard links publish without replacement even if another process races.
        published = []
        try:
            for name in names:
                os.link(temporary / name, output_dir / name)
                published.append(output_dir / name)
        except Exception:
            for path in published:
                path.unlink()
            raise
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exp1-root', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--mother-docx', type=Path)
    parser.add_argument('--allow-missing-mother', action='store_true')
    args = parser.parse_args()
    result = render_reports(args.exp1_root, args.output_dir, args.mother_docx, args.allow_missing_mother)
    print(json.dumps({'output_dir': str(args.output_dir.resolve()), 'excel_sheets': len(result['sheets']),
                      'word_summary_tables': len(result['word_tables']),
                      'mother_found': result['mother_document']['found']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
