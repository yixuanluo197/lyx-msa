#!/usr/bin/env python3
"""Independently recompute experiment 6 from original CSVs using NumPy."""
import argparse
import csv
import hashlib
import json
import math
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import numpy as np


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as stream:
        return list(csv.DictReader(stream))


def sha256(path):
    digest=hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda:stream.read(4*1024**2),b''):
            digest.update(block)
    return digest.hexdigest()


def check(value,expected):
    if expected is None:
        assert value=='N/A',value
    else:
        assert math.isclose(float(value),expected,rel_tol=0,abs_tol=1e-12),(value,expected)


def verify_statistics(output):
    manifest=json.loads((output/'run_manifest.json').read_text())
    summaries=read_csv(output/'exp6_overall_summary.csv')
    bins=read_csv(output/'exp6_confidence_bins.csv')
    assert len(summaries)==4 and len(bins)==40
    references,metadata={},{}
    verified_rows=0
    for group,source in manifest['source_predictions'].items():
        dataset,model=group.split('/')
        assert sha256(source['path'])==source['sha256']
        rows=read_csv(source['path'])
        names=manifest['models'][group]['class_names']
        probabilities=np.array([[float(r['probability_'+name]) for name in names] for r in rows],dtype=np.float64)
        labels=np.array([int(r['true_label']) for r in rows])
        predicted=np.array([int(r['predicted_label']) for r in rows])
        ids=[(r['dataset'],r['dialogue_id'],int(r['utterance_index'])) for r in rows]
        assert len(set(ids))==len(rows)==source['rows']
        assert len(rows)=={'IEMOCAP':1623,'MELD':2610}[dataset]
        assert np.isfinite(probabilities).all() and (probabilities>=0).all() and (probabilities<=1).all()
        assert np.allclose(probabilities.sum(1),1,rtol=0,atol=1e-5)
        assert np.array_equal(probabilities.argmax(1),predicted)
        assert np.allclose(probabilities.max(1),[float(r['max_probability']) for r in rows],rtol=0,atol=1e-12)
        true_probability=probabilities[np.arange(len(rows)),labels]
        assert (true_probability>0).all()
        assert np.allclose(true_probability,[float(r['true_class_probability']) for r in rows],rtol=0,atol=1e-12)
        confidence=probabilities.max(1)
        correct=(predicted==labels)
        edges=np.array([i/10 for i in range(11)])
        membership=np.minimum(np.searchsorted(edges,confidence,side='right')-1,9)
        stat=next(r for r in summaries if (r['dataset'],r['model'])==(dataset,model))
        detail=sorted((r for r in bins if (r['dataset'],r['model'])==(dataset,model)),key=lambda r:int(r['bin_index']))
        assert [int(r['bin_index']) for r in detail]==list(range(10))
        ece,maximum=0.0,0.0
        for index,row in enumerate(detail):
            selected=membership==index
            count=int(selected.sum())
            check(row['count'],count)
            check(row['lower_bound'],index/10)
            check(row['upper_bound'],(index+1)/10)
            check(row['weight'],count/len(rows))
            if count:
                mean=float(confidence[selected].mean());accuracy=float(correct[selected].mean())
                difference=mean-accuracy
                check(row['mean_confidence'],mean);check(row['accuracy'],accuracy)
                check(row['signed_gap'],difference);check(row['absolute_gap'],abs(difference))
                contribution=count/len(rows)*abs(difference)
                check(row['ece_contribution'],contribution)
                ece+=contribution;maximum=max(maximum,abs(difference))
            else:
                for field in ('mean_confidence','accuracy','signed_gap','absolute_gap'):check(row[field],None)
                check(row['ece_contribution'],0)
        assert sum(int(r['count']) for r in detail)==len(rows)
        targets=np.eye(len(names))[labels]
        expected=dict(valid_utterances=len(rows),accuracy=float(correct.mean()),mean_confidence=float(confidence.mean()),
                      ece=ece,mce=maximum,brier_score=float(np.square(probabilities-targets).sum(1).mean()),
                      nll=float(-np.log(true_probability).mean()),n_bins=10,zero_true_probability_count=0)
        for field,value in expected.items():check(stat[field],value)
        references[group]=expected
        metadata[dataset,model]=[(r['dialogue_id'],r['utterance_index'],r['source_utterance_id'],r['text'],r['true_label']) for r in rows]
        verified_rows+=len(rows)
    for dataset in ('IEMOCAP','MELD'):
        assert metadata[dataset,'baseline']==metadata[dataset,'cmerc']
        assert read_csv(output/(dataset+'_confidence_bins.csv'))==[r for r in bins if r['dataset']==dataset]
    for path,expected in manifest['previous_inputs_before_sha256'].items():
        assert sha256(path)==expected,path
    return dict(status='passed',independent_method='NumPy recomputation from original full-input CSVs; production metric implementation not imported',
                samples_checked=verified_rows,summary_rows_checked=4,bin_rows_checked=40,
                original_artifacts_hashes_verified=len(manifest['previous_inputs_before_sha256']),
                labels_text_ids_probabilities_aligned=True,empty_bins_na_verified=True,
                references=references)


def verify_reports(output):
    from docx import Document
    from lxml import etree

    manifest=json.loads((output/'render_manifest.json').read_text())
    for entry in manifest['outputs'].values():
        assert sha256(entry['path'])==entry['sha256'],entry['path']
    for entry in manifest['sources'].values():
        assert sha256(entry['path'])==entry['sha256'],entry['path']
    mother=Path(manifest['mother_document']['path'])
    assert sha256(mother)==manifest['mother_document']['sha256']
    document=Document(output/'CMERC_高置信度错误诊断报告_实验一至六.docx')
    original=Document(mother)
    def body_prefix(doc):
        return [etree.tostring(child,method='c14n') for child in doc.element.body
                if etree.QName(child).localname!='sectPr']
    assert body_prefix(document)[:len(body_prefix(original))]==body_prefix(original)
    assert etree.tostring(document.element.body.sectPr,method='c14n')==etree.tostring(original.element.body.sectPr,method='c14n')
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    val='{'+ns['w']+'}val'
    for entry in manifest['word_tables']:
        table=document.tables[entry['document_table_index']]
        actual=[[cell.text for cell in row.cells] for row in table.rows]
        source=read_csv(entry['source_csv'])
        expected=[]
        for index in entry['source_row_indices']:
            row=source[index]
            if entry['kind']=='overall':
                expected.append([row['dataset'],{'baseline':'Baseline','cmerc':'CMERC'}[row['model']],
                                 f"{100*float(row['ece']):.4f}",f"{100*float(row['mce']):.4f}",
                                 f"{float(row['brier_score']):.6f}",f"{float(row['nll']):.6f}"])
            else:
                interval=f"[{float(row['lower_bound']):.1f}, {float(row['upper_bound']):.1f}{']' if int(row['bin_index'])==9 else ')'}"
                expected.append([interval,row['count'],
                                 'N/A' if row['mean_confidence']=='N/A' else f"{100*float(row['mean_confidence']):.4f}",
                                 'N/A' if row['accuracy']=='N/A' else f"{100*float(row['accuracy']):.4f}",
                                 'N/A' if row['signed_gap']=='N/A' else f"{100*float(row['signed_gap']):+.6f}"])
        assert actual[0]==entry['headers'] and actual[1:]==expected
        borders={etree.QName(edge).localname:edge.get(val) for edge in table._tbl.tblPr.find('w:tblBorders',ns)}
        assert borders==dict(top='single',left='nil',bottom='single',right='nil',insideH='nil',insideV='nil')
        for cell in table.rows[0].cells:
            assert cell._tc.tcPr.find('w:tcBorders/w:bottom',ns).get(val)=='single'
    for entry in manifest['figures']:
        assert entry['axes']==dict(xlim=[0,1],ylim=[0,1],aspect='equal')
        for curve in entry['curves']:
            source=read_csv(curve['source_csv'])
            expected_points=[]
            expected_counts=[]
            selected=[r for r in source if r['dataset']==entry['dataset'] and r['model']==curve['model']]
            assert len(selected)==10
            for row in selected:
                expected_counts.append(int(row['count']))
                if int(row['count']):expected_points.append([float(row['mean_confidence']),float(row['accuracy'])])
            assert curve['points']==expected_points and curve['counts']==expected_counts
        assert Path(entry['files']['pdf']['path']).read_bytes().startswith(b'%PDF')
        svg=Path(entry['files']['svg']['path']).read_text()
        assert '<text' in svg,'SVG text must remain editable'
    spreadsheet=output/'CMERC_实验六_概率校准结果.xlsx'
    xml_ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(spreadsheet) as archive:
        assert archive.testzip() is None
        strings=[''.join(node.itertext()) for node in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
        for sheet in manifest['sheets']:
            data=ET.fromstring(archive.read(f"xl/worksheets/sheet{sheet['worksheet_index']}.xml"))
            cells=[]
            for row in data.findall('s:sheetData/s:row',xml_ns):
                values=[]
                for cell in row.findall('s:c',xml_ns):
                    assert cell.find('s:f',xml_ns) is None
                    raw=cell.findtext('s:v',namespaces=xml_ns)
                    values.append(strings[int(raw)] if cell.attrib.get('t')=='s' else raw)
                cells.append(values)
            source=read_csv(sheet['source_csv'])
            expected=[[source[i][column] for column in sheet['columns']] for i in sheet['source_row_indices']]
            assert cells[0]==sheet['columns'] and cells[1:]==expected
    with ZipFile(output/'CMERC_高置信度错误诊断报告_实验一至六.docx') as archive:
        assert archive.testzip() is None
    return dict(word_tables_checked=len(manifest['word_tables']),three_line_format_verified=True,
                original_word_body_and_section_unchanged=True,excel_sheets_checked=len(manifest['sheets']),
                all_word_excel_cells_equal_csv=True,reliability_figures_checked=len(manifest['figures']),
                plotted_points_and_histograms_equal_csv=True,editable_svg_text_verified=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,required=True)
    parser.add_argument('--statistics-only',action='store_true')
    options=parser.parse_args()
    output=options.output_dir.resolve()
    result=verify_statistics(output)
    if not options.statistics_only:result.update(verify_reports(output))
    print(json.dumps(result,ensure_ascii=False,indent=2))
