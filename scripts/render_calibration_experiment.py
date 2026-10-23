"""Append calibration diagnostics to a preserved report using fixed CSV inputs.

This renderer performs no training, inference, probability adjustment or model
selection. It uses Python/matplotlib exclusively for calibration figures.
"""
import argparse
import csv
import json
import math
import os
from pathlib import Path
import tempfile

import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import xlsxwriter

try:
    from .render_high_confidence_followups import (
        sha256, string_hash, body_prefix, display, validate_text,
        _append_markdown, _narrative_markdown)
except ImportError:
    from render_high_confidence_followups import (
        sha256, string_hash, body_prefix, display, validate_text,
        _append_markdown, _narrative_markdown)

DOCX_NAME = 'CMERC_高置信度错误诊断报告_实验一至六.docx'
XLSX_NAME = 'CMERC_实验六_概率校准结果.xlsx'
MANIFEST_NAME = 'render_manifest.json'
DATASETS = ('IEMOCAP', 'MELD')
MODELS = ('baseline', 'cmerc')
INPUT_FILES = ('exp6_overall_summary.csv', 'exp6_confidence_bins.csv', 'SUMMARY.md', 'run_manifest.json')
COLORS = {'baseline': '#377EB8', 'cmerc': '#E68632'}
LABELS = {'baseline': 'Baseline', 'cmerc': 'CMERC'}
CAPTION = ('可靠性图仅绘制非空箱的平均最大类别概率与实际准确率；虚线为理想校准。'
           '下图为全部十个区间的样本数，空箱不虚构为准确率为零的点。'
           '小样本箱估计不稳定；结果是当前本地 validation 权重的单 checkpoint 描述性比较，'
           '不支持显著性、跨 seed 稳健性或模态冲突的因果结论。')


def _read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        headers, rows = reader.fieldnames, list(reader)
    if not headers or len(headers) != len(set(headers)):
        raise ValueError(f'Invalid CSV headers: {path}')
    for row in rows:
        if None in row or any(v is None for v in row.values()):
            raise ValueError(f'Malformed CSV row: {path}')
        for value in row.values():
            validate_text(value, str(path))
    return headers, rows


def _number(value, name, allow_inf=False):
    result = float(value)
    if math.isnan(result) or (not allow_inf and not math.isfinite(result)):
        raise ValueError(f'Invalid numeric {name}: {value}')
    return result


def _missing(value):
    return value in ('', 'None', 'N/A', 'null', 'nan', 'NaN')


def _close(left, right):
    return math.isclose(left, right, rel_tol=1e-7, abs_tol=1e-8)


def _load(output_dir):
    for name in INPUT_FILES:
        if not (output_dir/name).is_file():
            raise FileNotFoundError(output_dir/name)
    sh, summaries = _read_csv(output_dir/INPUT_FILES[0])
    bh, bins = _read_csv(output_dir/INPUT_FILES[1])
    required_summary = {'dataset','model','input_condition','valid_utterances','accuracy','mean_confidence',
                        'ece','mce','brier_score','nll','n_bins','zero_true_probability_count'}
    required_bins = {'dataset','model','bin_index','lower_bound','upper_bound','count',
                     'mean_confidence','accuracy','signed_gap','absolute_gap','weight','ece_contribution'}
    if not required_summary.issubset(sh) or not required_bins.issubset(bh):
        raise ValueError('Missing calibration source columns')
    expected = {(d,m) for d in DATASETS for m in MODELS}
    groups = {}
    for index, row in enumerate(summaries):
        key = row['dataset'], row['model']
        if key not in expected or key in groups or row['input_condition'] != 'full':
            raise ValueError('Duplicate/unknown calibration model or non-full input')
        count = int(row['valid_utterances'])
        if count <= 0 or int(row['n_bins']) != 10:
            raise ValueError('Calibration requires positive support and ten bins')
        groups[key] = dict(summary=row, summary_row=index, bins=[])
    if set(groups) != expected:
        raise ValueError('Expected both datasets and both models in calibration summaries')
    for index, row in enumerate(bins):
        key = row['dataset'], row['model']
        if key not in groups:
            raise ValueError('Unknown bin dataset/model')
        groups[key]['bins'].append((index,row))
    for key, group in groups.items():
        records = sorted(group['bins'], key=lambda pair: int(pair[1]['bin_index']))
        if [int(row['bin_index']) for _,row in records] != list(range(10)):
            raise ValueError(f'Missing or duplicate bins: {key}')
        group['bins'] = records
        support = int(group['summary']['valid_utterances'])
        total, weighted_acc, weighted_conf, ece, gaps = 0,0,0,0,[]
        for _, row in records:
            index, count = int(row['bin_index']), int(row['count'])
            lo, hi = _number(row['lower_bound'],'lower_bound'), _number(row['upper_bound'],'upper_bound')
            if count < 0 or not _close(lo,index/10) or not _close(hi,(index+1)/10):
                raise ValueError('Invalid bin support or fixed ten-bin boundaries')
            if not _close(_number(row['weight'],'weight'), count/support):
                raise ValueError('Bin weight does not match denominator')
            total += count
            if not count:
                if any(not _missing(row[f]) for f in ('mean_confidence','accuracy','signed_gap','absolute_gap')):
                    raise ValueError('An empty bin must have undefined mean/accuracy/gaps')
                if not _close(_number(row['ece_contribution'],'ece_contribution'),0):
                    raise ValueError('Empty bin contributes nonzero ECE')
                continue
            confidence, accuracy = _number(row['mean_confidence'],'mean_confidence'), _number(row['accuracy'],'accuracy')
            if not (0 <= accuracy <= 1 and lo-1e-10 <= confidence <= hi+1e-10):
                raise ValueError('Bin mean outside interval or invalid accuracy')
            gap = confidence-accuracy
            if not _close(_number(row['signed_gap'],'signed_gap'),gap) or not _close(_number(row['absolute_gap'],'absolute_gap'),abs(gap)):
                raise ValueError('Inconsistent signed or absolute calibration gap')
            contribution = count/support*abs(gap)
            if not _close(_number(row['ece_contribution'],'ece_contribution'),contribution):
                raise ValueError('Inconsistent bin ECE contribution')
            ece += contribution
            gaps.append(abs(gap))
            weighted_acc += count/support*accuracy
            weighted_conf += count/support*confidence
        summary = group['summary']
        if total != support or any(not _close(_number(summary[field],field),expected_value) for field,expected_value in (
                ('accuracy',weighted_acc),('mean_confidence',weighted_conf),('ece',ece),('mce',max(gaps)))):
            raise ValueError(f'Summary does not match bin accounting: {key}')
        brier = _number(summary['brier_score'],'brier_score')
        nll = _number(summary['nll'],'nll',allow_inf=True)
        zero = int(summary['zero_true_probability_count'])
        if (not 0 <= brier <= 2 or nll < 0 or not 0 <= zero <= support
                or (not zero and not math.isfinite(nll)) or (zero and not math.isinf(nll))):
            raise ValueError('Invalid Brier/NLL or zero-probability accounting')
    return (sh,summaries), (bh,bins), groups


def _contract(groups):
    counts = '; '.join(f'{d}/{m}: n={g["summary"]["valid_utterances"]}, '
                       f'nonempty bins={sum(int(r["count"])>0 for _,r in g["bins"])}'
                       for (d,m),g in groups.items())
    return ('# Figure contract — Experiment 6\n\n'
            'Backend: Python/matplotlib only. Archetype: quantitative grid.\n\n'
            'Claim: compare observed confidence–accuracy gaps of these fixed Baseline and CMERC checkpoints; '
            'the figure alone makes no universal improvement or significance claim.\n\n'
            'Evidence: upper panel plots nonempty-bin mean maximum-class probability against empirical accuracy '
            'with equal 0–1 axes and an ideal diagonal; lower panel reports all ten bin counts. '
            'Empty bins have no reliability point. Connected points are guides, not fitted estimates.\n\n'
            f'Support: {counts}.\n\n'
            'Exports: 4.6 × 5.8 inch, PNG 300 dpi for Word; editable-text SVG and TrueType PDF. '
            'Baseline uses blue circles, CMERC orange squares; the legend is outside the axes.\n\n'
            'Review risks: sparse bins are unstable; no confidence intervals or paired significance are estimated. '
            'This is descriptive analysis of the current local validation weights, one checkpoint per model. '
            'No retraining, checkpoint selection, probability clipping or calibration fitting is performed.\n')


def _figures(directory, final_dir, groups):
    figures = []
    with plt.rc_context({'font.family':'DejaVu Sans','font.size':8,'axes.labelsize':8,
                         'axes.titlesize':9,'xtick.labelsize':7,'ytick.labelsize':7,
                         'legend.fontsize':8,'svg.fonttype':'none','pdf.fonttype':42,
                         'axes.spines.top':False,'axes.spines.right':False}):
        for dataset in DATASETS:
            fig, (ax, count_ax) = plt.subplots(2,1,figsize=(4.6,5.8),gridspec_kw={'height_ratios':[3,1]})
            fig.subplots_adjust(left=.18,right=.95,top=.87,bottom=.10,hspace=.42)
            ax.plot([0,1],[0,1],color='#777777',linestyle='--',linewidth=.9,label='_ideal')
            curves, handles = [], []
            for offset, model in zip((-.020,.020),MODELS):
                group = groups[dataset,model]
                active = [(i,r) for i,r in group['bins'] if int(r['count'])>0]
                points = [[float(r['mean_confidence']),float(r['accuracy'])] for _,r in active]
                line, = ax.plot([p[0] for p in points],[p[1] for p in points],
                                color=COLORS[model],marker='o' if model=='baseline' else 's',
                                markersize=4,linewidth=1.2,label=LABELS[model])
                handles.append(line)
                counts = [int(r['count']) for _,r in group['bins']]
                centers = [(float(r['lower_bound'])+float(r['upper_bound']))/2+offset for _,r in group['bins']]
                count_ax.bar(centers,counts,width=.036,color=COLORS[model],alpha=.85)
                curves.append(dict(model=model,source_csv=str(final_dir/'exp6_confidence_bins.csv'),
                                   bin_indices=[int(r['bin_index']) for _,r in active],
                                   source_row_indices=[i for i,_ in active],points=points,counts=counts,
                                   count_source_row_indices=[i for i,_ in group['bins']]))
            ax.set(xlim=(0,1),ylim=(0,1),xlabel='Mean confidence',ylabel='Empirical accuracy',title=dataset)
            ax.set_aspect('equal',adjustable='box')
            ax.set_xticks([0,.2,.4,.6,.8,1])
            ax.set_yticks([0,.2,.4,.6,.8,1])
            count_ax.set(xlim=(0,1),xlabel='Confidence bin',ylabel='Count')
            count_ax.set_xticks([0,.2,.4,.6,.8,1])
            fig.legend(handles,[LABELS[m] for m in MODELS],loc='upper center',bbox_to_anchor=(.56,.98),ncol=2,frameon=False)
            files = {}
            for extension in ('png','pdf','svg'):
                name = f'figures/reliability_{dataset}.{extension}'
                fig.savefig(directory/name,dpi=300,facecolor='white')
                files[extension] = dict(path=str(final_dir/name),sha256=sha256(directory/name))
            plt.close(fig)
            point_map = {curve['model']:[dict(bin_index=bin_index,count=curve['counts'][bin_index],
                                             source_row_index=source_index,x=point[0],y=point[1])
                                        for bin_index,source_index,point in zip(curve['bin_indices'],curve['source_row_indices'],curve['points'])]
                         for curve in curves}
            figures.append(dict(dataset=dataset,files=files,paths={fmt:entry['path'] for fmt,entry in files.items()},
                                curves=curves,points=point_map,caption=CAPTION,
                                axes=dict(xlim=[0,1],ylim=[0,1],aspect='equal'),png_dpi=300))
    return figures


def _three_line_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'),'single' if side in ('top','bottom') else 'nil')
        edge.set(qn('w:sz'),'8')
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        properties = cell._tc.get_or_add_tcPr()
        edges, edge = OxmlElement('w:tcBorders'), OxmlElement('w:bottom')
        edge.set(qn('w:val'),'single')
        edge.set(qn('w:sz'),'4')
        edges.append(edge)
        properties.append(edges)
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(table.add_row().cells,row):
            cell.text = text
    for row_index,row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    run.bold = row_index == 0
    return table


def _formatted(value, percent=False, signed=False):
    if _missing(value):
        return 'N/A'
    number = float(value)*(100 if percent else 1)
    if math.isinf(number):
        return '+∞' if number>0 else '−∞'
    return format(number, '+.6f' if signed else '.4f' if percent else '.6f')


def _word(directory, final_dir, mother, groups, figures):
    original = Document(mother)
    prefix, mother_hash = body_prefix(original), sha256(mother)
    # Append into the existing final section. Do not rewrite its dimensions,
    # terminal sectPr, or any original section/page/style configuration.
    original.add_page_break()
    original.add_heading('实验六：Baseline 与 CMERC 的置信度校准诊断',1)
    original.add_paragraph('沿用固定 checkpoint 和完整文本、音频、视觉输入，复用已有测试集逐话语概率，不重训练、不改概率或标签，不以测试集选模或调参。')
    original.add_heading('指标定义与统计口径',2)
    original.add_paragraph('预测置信度 cᵢ=maxₖpᵢₖ；正确性 aᵢ=1[argmaxₖpᵢₖ=yᵢ]。使用十个等宽区间 [0,0.1)、…、[0.9,1]，排除 padding。对非空箱 b：conf(b)=Σcᵢ/n_b，acc(b)=Σaᵢ/n_b，gap(b)=conf(b)−acc(b)。正差表示过置信，负差表示欠置信。')
    original.add_paragraph('ECE=Σ_b(n_b/N)|gap(b)|；MCE=max_{b:n_b>0}|gap(b)|。空箱的准确率、平均置信度和差值为 N/A，不纳入 MCE。ECE/MCE 在下表乘以100并标为 %。')
    original.add_paragraph('多分类 Brier=(1/N)ΣᵢΣₖ(pᵢₖ−1[yᵢ=k])²，不除以类别数、不除以2，范围0–2。NLL=−(1/N)Σᵢln pᵢ,yᵢ，使用自然对数，不裁剪概率；真实类别概率为零时理论值为+∞，须同时检查清单中的零概率计数。四项校准/概率评分均越低越好，但不能代替分类准确率或 W-F1。')
    tables = []
    def add_table(kind, headers, rows, source, indices, columns, dataset=None, model=None):
        index = len(original.tables)
        _three_line_table(original,headers,rows)
        tables.append(dict(kind=kind,dataset=dataset,model=model,source_csv=str(final_dir/source),
                           source_row_indices=indices,source_columns=columns,
                           document_table_index=index,headers=headers,columns=headers,row_count=len(rows),
                           rendered_rows=rows,rendered_rows_sha256=string_hash(rows),
                           display_rule=('model mapped Baseline/CMERC; ece and mce multiplied by 100, 4 decimals; '
                                         'Brier/NLL 6 decimals; positive infinity written +∞' if kind=='overall' else
                                         'bounds formatted 1 decimal, final interval closed; count unchanged; '
                                         'mean confidence and accuracy multiplied by 100, 4 decimals (%); '
                                         'signed gap multiplied by 100, explicitly signed 6 decimals (pp); empty bins N/A'),
                           cells_sha256=string_hash([headers]+rows)))
    original.add_heading('总体指标对比',2)
    headers = ['Dataset','Model','ECE (%)','MCE (%)','Brier','NLL']
    rows, indices = [], []
    for dataset in DATASETS:
        for model in MODELS:
            group = groups[dataset,model]
            row = group['summary']
            rows.append([dataset,LABELS[model],_formatted(row['ece'],True),_formatted(row['mce'],True),
                         _formatted(row['brier_score']),_formatted(row['nll'])])
            indices.append(group['summary_row'])
    add_table('overall',headers,rows,'exp6_overall_summary.csv',indices,
              ['dataset','model','ece','mce','brier_score','nll'])
    original.add_paragraph('表注：ECE、MCE 使用百分数；Brier、NLL 使用原始分值。该表只比较清单中的当前权重，不能据此挑选 checkpoint。有效话语数、准确率、平均置信度及零真实类概率计数完整保存在 CSV 和独立 Excel。')
    for dataset in DATASETS:
        original.add_heading(f'{dataset}：十箱统计与可靠性图',2)
        for model in MODELS:
            original.add_heading(LABELS[model],3)
            group = groups[dataset,model]
            headers = ['Confidence interval','Count','Mean confidence (%)','Accuracy (%)','Signed gap (pp)']
            rows = []
            for _,row in group['bins']:
                lo,hi,index = float(row['lower_bound']),float(row['upper_bound']),int(row['bin_index'])
                interval = f'[{lo:.1f}, {hi:.1f}{"]" if index==9 else ")"}'
                rows.append([interval,row['count'],_formatted(row['mean_confidence'],percent=True),
                             _formatted(row['accuracy'],percent=True),_formatted(row['signed_gap'],percent=True,signed=True)])
            add_table('bins',headers,rows,'exp6_confidence_bins.csv',[i for i,_ in group['bins']],
                      ['lower_bound','upper_bound','count','mean_confidence','accuracy','signed_gap'],dataset,model)
            original.add_paragraph('表注：本表平均置信度和准确率以百分数 (%) 显示，差值以百分点 (pp) 显示，均由源 CSV 的0–1值乘以100。正差为过置信、负差为欠置信。各箱 count 之和等于该模型有效测试话语总数；N/A 是空箱的未定义统计量。源 CSV、Excel 和图坐标仍保留0–1尺度。')
        figure = next(f for f in figures if f['dataset']==dataset)
        original.add_picture(str(directory/'figures'/Path(figure['files']['png']['path']).name),width=Inches(4.15))
        original.add_paragraph(f'{dataset} Reliability Diagram。{CAPTION}')
    original.add_heading('实际发现、局限与后续工作',2)
    narrative = (final_dir/'SUMMARY.md').read_text(encoding='utf-8')
    validate_text(narrative,'SUMMARY.md')
    _append_markdown(original,_narrative_markdown(narrative))
    if body_prefix(original)[:len(prefix)] != prefix:
        raise AssertionError('Original Word body changed before saving')
    original.save(directory/DOCX_NAME)
    saved = Document(directory/DOCX_NAME)
    if body_prefix(saved)[:len(prefix)] != prefix or sha256(mother) != mother_hash:
        raise AssertionError('Original report content or original file changed')
    for table in tables:
        cells = [[cell.text for cell in row.cells] for row in saved.tables[table['document_table_index']].rows]
        if string_hash(cells) != table['cells_sha256']:
            raise AssertionError('Rendered calibration table differs after serialization')
    return tables, dict(path=str(mother),sha256=mother_hash,body_prefix_preserved=True,
                        body_element_count=len(prefix),body_sha256=string_hash([x.decode('utf-8') for x in prefix]))


def _excel(path, summaries, bins, groups, final_dir):
    sheets = []
    with xlsxwriter.Workbook(str(path),{'strings_to_formulas':False,'strings_to_urls':False}) as workbook:
        head = workbook.add_format({'bold':True,'bg_color':'#DDEBF7','text_wrap':True})
        body = workbook.add_format({'text_wrap':True,'valign':'top'})
        def sheet(name, headers, rows, source, indices):
            work = workbook.add_worksheet(name)
            work.write_row(0,0,headers,head)
            rendered = []
            for index,row in enumerate(rows,1):
                values = [display(row[key]) for key in headers]
                rendered.append(values)
                for column,value in enumerate(values):
                    if work.write_string(index,column,value,body) != 0:
                        raise ValueError('Excel could not preserve a source cell')
            work.freeze_panes(1,0)
            work.autofilter(0,0,len(rows),len(headers)-1)
            for index,header in enumerate(headers):
                work.set_column(index,index,min(32,max(15,len(header)+2)))
            sheets.append(dict(name=name,worksheet_index=len(sheets)+1,source_csv=str(final_dir/source),source_row_indices=indices,
                               row_count=len(rows),columns=headers,rendered_rows_sha256=string_hash(rendered)))
        sheet('Overall',*summaries,'exp6_overall_summary.csv',list(range(len(summaries[1]))))
        for dataset in DATASETS:
            for model in MODELS:
                group = groups[dataset,model]
                sheet(f'{dataset}_{model}',bins[0],[row for _,row in group['bins']],
                      'exp6_confidence_bins.csv',[i for i,_ in group['bins']])
    return sheets


def render_calibration_reports(output_dir,mother_docx):
    output_dir, mother = Path(output_dir).resolve(), Path(mother_docx).resolve()
    if not mother.is_file():
        raise FileNotFoundError(mother)
    names = [DOCX_NAME,XLSX_NAME,MANIFEST_NAME,'figure_contract.md'] + [
        f'figures/reliability_{dataset}.{extension}' for dataset in DATASETS for extension in ('png','pdf','svg')]
    for name in names:
        if (output_dir/name).exists():
            raise FileExistsError(f'Refusing to overwrite calibration report: {output_dir/name}')
    summaries,bins,groups = _load(output_dir)
    source_hashes = {name:sha256(output_dir/name) for name in INPUT_FILES}
    with tempfile.TemporaryDirectory(prefix='.render_calibration_',dir=output_dir) as folder:
        temporary = Path(folder)
        # The figure contract exists before any graphics device is created.
        (temporary/'figure_contract.md').write_text(_contract(groups),encoding='utf-8')
        (temporary/'figures').mkdir()
        figures = _figures(temporary,output_dir,groups)
        tables,provenance = _word(temporary,output_dir,mother,groups,figures)
        sheets = _excel(temporary/XLSX_NAME,summaries,bins,groups,output_dir)
        manifest = dict(schema_version=1,mother_document=provenance,word_tables=tables,
                        figures=figures,sheets=sheets,sources={name:dict(path=str(output_dir/name),sha256=digest)
                                                             for name,digest in source_hashes.items()},
                        outputs={name:dict(path=str(output_dir/name),sha256=sha256(temporary/name))
                                 for name in names if name!=MANIFEST_NAME})
        for name,digest in source_hashes.items():
            if sha256(output_dir/name) != digest:
                raise RuntimeError(f'Calibration input changed during rendering: {name}')
        (temporary/MANIFEST_NAME).write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        installed = []
        try:
            for name in names:
                (output_dir/name).parent.mkdir(parents=True,exist_ok=True)
                os.link(temporary/name,output_dir/name)
                installed.append(output_dir/name)
        except Exception:
            for path in installed:
                path.unlink()
            raise
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,required=True)
    parser.add_argument('--mother-docx',type=Path,required=True)
    args = parser.parse_args()
    manifest = render_calibration_reports(args.output_dir,args.mother_docx)
    print(json.dumps({'word_tables':len(manifest['word_tables']),'figures':len(manifest['figures']),
                      'excel_sheets':len(manifest['sheets'])},ensure_ascii=False))


if __name__ == '__main__':
    main()
