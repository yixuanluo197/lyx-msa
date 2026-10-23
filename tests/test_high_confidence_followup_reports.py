"""Report rendering checks using synthetic data, never experiment artifacts."""
import csv
import hashlib
import importlib
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document


SUMMARY_FILES = (
    'exp2_summary.csv', 'exp3_class_summary.csv',
    'exp4_modality_ablation_summary.csv', 'exp4_full_correct_to_masked_error.csv',
    'exp5_error_confidence_distribution.csv', 'exp5_overall_summary.csv',
)


def renderer():
    return importlib.import_module('scripts.render_high_confidence_followups')


def fixtures(tmp_path):
    root = tmp_path / 'exp1'
    out = root / 'experiments_2_to_5'
    out.mkdir(parents=True)
    (root / 'SUMMARY.md').write_text('# 实验一\n\n原始内容不可更改。\n\n| x | y |\n|---|---|\n| 1 | 2 |\n', encoding='utf-8')
    for name in (*SUMMARY_FILES, 'exp2_paired_outcomes.csv'):
        with (out / name).open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(['dataset', 'metric', 'value', 'utterance_text'])
            writer.writerow(['IEMOCAP', 'errors', 'None', '=SUM(1,2)'])
            writer.writerow(['MELD', 'errors', 'N/A', 'https://example.test/path'])
            writer.writerow(['IEMOCAP', 'correct', '3', 'line1\nline2\tend'])
    return root, out


def test_missing_mother_requires_explicit_fallback(tmp_path):
    root, out = fixtures(tmp_path)
    with pytest.raises(FileNotFoundError, match='母版'):
        renderer().render_reports(root, out)
    assert not list(out.glob('*.docx'))
    manifest = renderer().render_reports(root, out, allow_missing_mother=True)
    assert manifest['mother_document']['found'] is False
    document = Document(out / renderer().DOCX_NAME)
    assert '未找到指定原Word母版' in '\n'.join(p.text for p in document.paragraphs)
    assert any('原始内容不可更改。' in p.text for p in document.paragraphs)
    assert document.tables[0].cell(1, 0).text == '1'


def test_mother_body_and_file_remain_unchanged(tmp_path):
    root, out = fixtures(tmp_path)
    mother = root / 'CMERC_实验一_高置信度错误样本统计表.docx'
    document = Document()
    document.add_paragraph('保留原段落', style='Title')
    document.add_table(rows=1, cols=1).cell(0, 0).text = '保留原表格'
    document.save(mother)
    before = mother.read_bytes()
    original = renderer().body_prefix(Document(mother))
    manifest = renderer().render_reports(root, out, mother)
    completed = Document(out / renderer().DOCX_NAME)
    assert renderer().body_prefix(completed)[:len(original)] == original
    assert mother.read_bytes() == before
    assert manifest['mother_document']['body_prefix_preserved'] is True
    assert manifest['mother_document']['sha256'] == hashlib.sha256(before).hexdigest()


def test_xlsx_full_rows_strings_na_and_dataset_split(tmp_path):
    root, out = fixtures(tmp_path)
    manifest = renderer().render_reports(root, out, allow_missing_mother=True)
    mappings = [s for s in manifest['sheets'] if s['source_csv'].endswith('exp2_paired_outcomes.csv')]
    assert {(s['dataset'], s['row_count']) for s in mappings} == {('IEMOCAP', 2), ('MELD', 1)}
    with ZipFile(out / renderer().XLSX_NAME) as archive:
        xml = archive.read('xl/sharedStrings.xml').decode('utf-8')
        assert '=SUM(1,2)' in xml and 'N/A' in xml
        assert 'https://example.test/path' in xml
        assert 'line1\nline2\tend' in xml
        for sheet in mappings:
            sheet_xml = archive.read(f'xl/worksheets/sheet{sheet["worksheet_index"]}.xml')
            element = ET.fromstring(sheet_xml)
            ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            assert len(element.findall('.//s:sheetData/s:row', ns)) == sheet['row_count'] + 1
            assert not element.findall('.//s:f', ns)
            assert len(sheet['name']) <= 31
    for table in manifest['word_tables']:
        with Path(table['source_csv']).open(encoding='utf-8', newline='') as stream:
            data = list(csv.DictReader(stream))
        rendered = [[renderer().display(data[i][h]) for h in table['columns']] for i in table['source_row_indices']]
        assert renderer().string_hash(rendered) == table['rendered_rows_sha256']


def test_existing_outputs_not_overwritten(tmp_path):
    root, out = fixtures(tmp_path)
    sentinel = out / renderer().XLSX_NAME
    sentinel.write_bytes(b'do not replace')
    with pytest.raises(FileExistsError):
        renderer().render_reports(root, out, allow_missing_mother=True)
    assert sentinel.read_bytes() == b'do not replace'
    assert not (out / renderer().DOCX_NAME).exists()


def test_control_characters_fail_explicitly_without_csv_edits(tmp_path):
    root, out = fixtures(tmp_path)
    source = out / 'exp2_paired_outcomes.csv'
    source.write_text('dataset,utterance_text\nIEMOCAP,abc\x01def\n', encoding='utf-8')
    before = source.read_bytes()
    with pytest.raises(ValueError, match='control'):
        renderer().render_reports(root, out, allow_missing_mother=True)
    assert source.read_bytes() == before


def test_wide_word_tables_keep_every_field_and_row_without_repeat_markdown(tmp_path):
    root, out = fixtures(tmp_path)
    source = out / 'exp5_overall_summary.csv'
    headers = ['dataset', 'model'] + [f'metric_{i}' for i in range(15)] + [
        f'{name}_ge_{threshold}' for threshold in ('0p8', '0p9', '0p95')
        for name in ('errors', 'high_confidence_predictions', 'high_confidence_error_rate')]
    with source.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        for model in ('baseline', 'cmerc'):
            writer.writerow({key: ('IEMOCAP' if key == 'dataset' else model if key == 'model' else '0.1234567890123456') for key in headers})
    original = (root / 'SUMMARY.md').read_text(encoding='utf-8')
    (out / 'SUMMARY.md').write_text(original+'\n## 主要发现\n实际观察文字。\n\n| repeat | table |\n|---|---|\n| not | duplicated |\n', encoding='utf-8')
    manifest = renderer().render_reports(root, out, allow_missing_mother=True)
    document = Document(out / renderer().DOCX_NAME)
    assert any(p.text == '实际观察文字。' for p in document.paragraphs)
    assert not any(t.cell(0, 0).text == 'repeat' for t in document.tables)
    groups = [t for t in manifest['word_tables'] if t['source_csv'] == str(source.resolve()) and t['dataset'] == 'IEMOCAP']
    assert len(groups) >= 3
    assert set(key for group in groups for key in group['columns']) == set(headers)
    for group in groups:
        assert len(group['columns']) <= 10
        assert {'dataset', 'model'} <= set(group['columns'])
        assert group['source_row_indices'] == [0, 1]
        actual = [[cell.text for cell in row.cells] for row in document.tables[group['document_table_index']].rows]
        assert renderer().string_hash(actual[1:]) == group['rendered_rows_sha256']


def test_c1_characters_are_preserved_and_empty_dataset_has_header_only_sheet(tmp_path):
    root, out = fixtures(tmp_path)
    source = out / 'exp2_paired_outcomes.csv'
    text = ''.join(chr(i) for i in range(0x80, 0xA0))
    with source.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['dataset', 'utterance_text'])
        writer.writerow(['MELD', text])
    before = source.read_bytes()
    manifest = renderer().render_reports(root, out, allow_missing_mother=True)
    assert source.read_bytes() == before
    sheets = [s for s in manifest['sheets'] if s['source_csv'] == str(source.resolve())]
    assert {(s['dataset'], s['row_count']) for s in sheets} == {('MELD', 1), ('IEMOCAP', 0)}
    with ZipFile(out / renderer().XLSX_NAME) as archive:
        shared = ET.fromstring(archive.read('xl/sharedStrings.xml'))
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        strings = [''.join(item.itertext()) for item in shared.findall('s:si', ns)]
        assert text in strings
        empty = next(s for s in sheets if s['dataset'] == 'IEMOCAP')
        sheet = ET.fromstring(archive.read(f'xl/worksheets/sheet{empty["worksheet_index"]}.xml'))
        assert len(sheet.findall('.//s:sheetData/s:row', ns)) == 1
